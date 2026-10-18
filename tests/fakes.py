"""Test doubles and builders. Nothing here touches the network."""

from __future__ import annotations

import random
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from statement_tracker.core.market_data.provider import (
    DailyRange,
    PriceUnavailableError,
    Quote,
    SymbolHistory,
    SymbolInfo,
    SymbolNotFoundError,
)
from statement_tracker.core.models import Statement, Transaction
from statement_tracker.core.parser import parse_rows

HEADER = "Date,Ticker,Type,Quantity,Price per share,Total Amount,Currency,FX Rate"
SAMPLES = Path(__file__).resolve().parents[1] / "samples"
SAMPLE = SAMPLES / "revolut_sample.csv"


def statement_from(*lines: str, header: str = HEADER) -> Statement:
    """Parse statement rows given as CSV lines (the header is added)."""
    rows = [(number, line.split(",")) for number, line in enumerate((header, *lines), start=1)]
    return parse_rows(rows)


def write_statement(folder: Path, *lines: str, header: str = HEADER) -> Path:
    """Write a statement CSV file and return its path."""
    path = folder / "statement.csv"
    path.write_text("\n".join((header, *lines)) + "\n", encoding="utf-8")
    return path


@dataclass
class Expected:
    """What a statement made by ``random_statement`` must add up to, kept by the generator as it writes
    each row: the plain average-cost rules of the specification (B1 to B7), worked out separately
    from the app's own calculator."""

    deposits: Decimal = Decimal(0)
    cash: Decimal = Decimal(0)
    shares: dict[str, Decimal] = field(default_factory=dict)
    cost_basis: dict[str, Decimal] = field(default_factory=dict)
    realized: dict[str, Decimal] = field(default_factory=dict)
    dividends: dict[str, Decimal] = field(default_factory=dict)

    @property
    def held(self) -> list[str]:
        """Tickers still held, in the order they were first bought."""
        return [ticker for ticker, shares in self.shares.items() if shares > 0]


def _amount(value: Decimal, currency: str, rng: random.Random) -> str:
    """An amount as Revolut writes it: "EUR 123.45", "EUR 1,234.56", or "EUR 120" when it is whole."""
    if value == value.to_integral_value() and rng.random() < 0.5:
        return f"{currency} {value:,.0f}" if rng.random() < 0.5 else f"{currency} {value:.0f}"
    return f"{currency} {value:,.2f}" if rng.random() < 0.7 else f"{currency} {value:.2f}"


def random_statement(
    seed: int,
    tickers: Sequence[str] = ("ABCD", "WXYZ", "QRST", "LMNO", "EFGH"),
    months: int = 36,
    start: datetime = datetime(2022, 1, 3, 8, 0),
    currency: str = "EUR",
) -> tuple[list[str], Expected]:
    """A believable full-history statement, different for each seed, as CSV lines (without the header),
    and what it must add up to.

    Every month there is a top-up, then buys of a few of the tickers (with a small fee in the total),
    sometimes a sale of part or all of a holding (minus a fee), and sometimes a dividend, also on
    tickers sold since. Buying again after selling everything happens too. Quantities have up to
    eight decimals; times have three or six fractional digits; amounts are written in each of the
    ways Revolut writes them. Nothing is ever sold that is not held.
    """
    import csv
    import io
    from datetime import timedelta

    rng = random.Random(seed)
    expected = Expected()
    price = {ticker: Decimal(rng.randint(2000, 50000)) / 100 for ticker in tickers}
    rows: list[list[str]] = []
    moment = start

    def at() -> str:
        nonlocal moment
        moment += timedelta(minutes=rng.randint(5, 600), seconds=rng.randint(0, 59))
        fraction = f"{rng.randint(0, 999):03d}" if rng.random() < 0.5 else f"{rng.randint(0, 999999):06d}"
        return moment.strftime("%Y-%m-%dT%H:%M:%S.") + fraction + "Z"

    for month in range(months):
        moment = max(moment, start + timedelta(days=30 * month))
        top_up = Decimal(rng.randint(100, 3000))
        rows.append([at(), "", "CASH TOP-UP", "", "", _amount(top_up, currency, rng), currency, "1"])
        expected.deposits += top_up
        expected.cash += top_up
        for ticker in tickers:
            price[ticker] = max(Decimal("1.00"), (price[ticker] * Decimal(rng.uniform(0.9, 1.12))).quantize(Decimal("0.01")))
        chosen = rng.sample(list(tickers), rng.randint(1, 3))
        budget = max(expected.cash, Decimal(0)) * Decimal(rng.uniform(0.4, 0.95)) / len(chosen)
        for ticker in chosen:  # most of the cash there is goes into the investments, as people do
            quantity = (budget / price[ticker]).quantize(Decimal("0.00000001"))
            if quantity <= 0:
                continue
            fee = Decimal(rng.choice(["0", "0", "0.01", "0.50", "1"]))
            total = (quantity * price[ticker] + fee).quantize(Decimal("0.01"))
            if total < Decimal("0.01"):
                continue
            rows.append([at(), ticker, "BUY - MARKET", str(quantity), _amount(price[ticker], currency, rng),
                         _amount(total, currency, rng), currency, "1"])  # fmt: skip
            expected.shares[ticker] = expected.shares.get(ticker, Decimal(0)) + quantity
            expected.cost_basis[ticker] = expected.cost_basis.get(ticker, Decimal(0)) + total
            expected.cash -= total
        for ticker in expected.held:
            if rng.random() < 0.12:
                held = expected.shares[ticker]
                everything = rng.random() < 0.35
                quantity = held if everything else (held * Decimal(rng.uniform(0.1, 0.8))).quantize(Decimal("0.00000001"))
                if not everything and held - quantity < Decimal("0.001"):
                    quantity = held  # never leave a crumb: sell all of it
                fee = Decimal(rng.choice(["0", "0.01", "1"]))
                total = max(Decimal("0.01"), (quantity * price[ticker] - fee).quantize(Decimal("0.01")))
                rows.append([at(), ticker, "SELL - MARKET", str(quantity), _amount(price[ticker], currency, rng),
                             _amount(total, currency, rng), currency, "1"])  # fmt: skip
                average = expected.cost_basis[ticker] / held
                expected.realized[ticker] = expected.realized.get(ticker, Decimal(0)) + total - average * quantity
                expected.shares[ticker] = held - quantity
                expected.cost_basis[ticker] = Decimal(0) if quantity == held else expected.cost_basis[ticker] - average * quantity
                expected.cash += total
        for ticker in list(expected.shares):  # held now or once
            if rng.random() < 0.08:
                dividend = Decimal(rng.randint(1, 2500)) / 100
                rows.append([at(), ticker, "DIVIDEND", "", "", _amount(dividend, currency, rng), currency, "1"])
                expected.dividends[ticker] = expected.dividends.get(ticker, Decimal(0)) + dividend
                expected.cash += dividend

    lines = []
    for row in rows:
        out = io.StringIO()
        csv.writer(out, lineterminator="").writerow(row)
        lines.append(out.getvalue())
    return lines, expected


def ranges_around(buys: Iterable[Transaction], spread: str = "0.50") -> dict[date, tuple[Decimal, Decimal]]:
    """Build a synthetic price history whose daily ranges contain the given BUY prices."""
    margin = Decimal(spread)
    return {
        buy.timestamp.date(): (buy.price_per_share - margin, buy.price_per_share + margin)
        for buy in buys
        if buy.price_per_share is not None
    }


class _Text(str):
    """A text cell: arithmetic on it is an error, as in Excel (#VALUE!)."""

    def _error(self, *_: object) -> _Text:
        return ERROR

    __add__ = __radd__ = __sub__ = __rsub__ = __mul__ = __rmul__ = __truediv__ = __rtruediv__ = _error


ERROR = _Text("#VALUE!")


def computed(path: Path, sheet_name: str) -> list[list[object]]:
    """Return a sheet's rows as Excel would show them, working out its formulas.

    Knows only what the app's sheets use: IF, OR, AND, N, ISNUMBER, COUNT, SUM, SUMPRODUCT, arithmetic,
    comparisons, cell references and ranges, also on other sheets (``$`` ignored). Numbers come back
    as floats.
    """
    import re

    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    workbook = load_workbook(path)
    names = list(workbook.sheetnames)
    done: dict[tuple[int, int, int], object] = {}

    def value(column: str, row: int, sheet: int) -> object:
        key = (sheet, row, column_index_from_string(column))
        if key not in done:
            raw = workbook[names[sheet]].cell(row=row, column=key[2]).value
            if isinstance(raw, str) and raw.startswith("="):
                done[key] = evaluate(raw[1:], sheet)
            elif isinstance(raw, (int, float, Decimal)) and not isinstance(raw, bool):
                done[key] = float(raw)
            else:
                done[key] = _Text(raw) if raw is not None else None
        return done[key]

    def cells(first: str, first_row: int, last: str, last_row: int, sheet: int) -> list[object]:
        return [value(first, row, sheet) for row in range(first_row, last_row + 1)]

    def sheet_index(match: re.Match[str]) -> str:
        return str(names.index(match.group(1).replace("''", "'")))

    def is_number(item: object) -> object:
        if isinstance(item, list):
            return [is_number(each) for each in item]
        return isinstance(item, float)

    functions = {
        "IF": lambda condition, then, otherwise: then if condition else otherwise,
        "OR": lambda *items: any(items),
        "AND": lambda *items: all(items),
        "N": lambda item: item if isinstance(item, float) else 0.0,
        "ISNUMBER": is_number,
        "COUNT": lambda items: sum(isinstance(item, float) for item in items),
        "SUM": lambda items: sum(item for item in items if isinstance(item, float)),
        "SUMPRODUCT": lambda left, right: sum(a * b for a, b in zip(left, right) if isinstance(a, float)),
        "_cells": cells,
        "_value": value,
    }

    def evaluate(formula: str, sheet: int) -> object:
        text = formula.replace("$", "").replace("--", "")
        # 'Other sheet'!B4 becomes «1»B4, the sheet's number, so that its name is never read as a cell.
        text = re.sub(r"'((?:[^']|'')+)'!", lambda match: f"«{sheet_index(match)}»", text)
        here = f"«{sheet}»"
        text = re.sub(r"(?<![»:\w])([A-Z]+\d+(?::[A-Z]+\d+)?)\b", lambda match: here + match.group(1), text)
        text = re.sub(r"«(\d+)»([A-Z]+)(\d+):([A-Z]+)(\d+)\b", r"_cells('\2',\3,'\4',\5,\1)", text)
        text = re.sub(r"«(\d+)»([A-Z]{1,2})(\d+)\b", r"_value('\2',\3,\1)", text)
        text = re.sub(r"(?<![<>!=])=(?!=)", "==", text)
        try:
            result = eval(text, {"__builtins__": {}}, functions)  # noqa: S307 - our own formulas, in tests
        except ZeroDivisionError:
            return ERROR
        if isinstance(result, (bool, int)):
            return float(result)
        return _Text(result) if isinstance(result, str) else result

    sheet = workbook[sheet_name]
    index = names.index(sheet_name)
    return [
        [value(chr(ord("A") + column), row, index) for column in range(sheet.max_column)]
        for row in range(1, sheet.max_row + 1)
    ]


@dataclass
class Listing:
    info: SymbolInfo
    price: Decimal | None
    ranges: dict[date, tuple[Decimal, Decimal]]
    as_of: datetime | None = None


@dataclass
class FakeProvider:
    """A PriceProvider with canned listings that records every call."""

    listings: dict[str, Listing] = field(default_factory=dict)
    search_results: dict[str, Sequence[SymbolInfo]] = field(default_factory=dict)
    offline: bool = False
    calls: list[tuple[str, str]] = field(default_factory=list)
    prepared: int = 0

    def add(
        self,
        symbol: str,
        currency: str,
        price: str | None = None,
        ranges: dict[date, tuple[Decimal, Decimal]] | None = None,
        name: str | None = None,
        exchange: str | None = None,
        timezone: str | None = "Europe/Berlin",
        as_of: datetime | None = None,
    ) -> None:
        info = SymbolInfo(symbol, name or f"{symbol} fund", exchange, currency, timezone)
        self.listings[symbol] = Listing(
            info, Decimal(price) if price is not None else None, ranges or {}, as_of
        )

    def prepare(self) -> None:
        self.prepared += 1

    def _listing(self, kind: str, symbol: str) -> Listing:
        self.calls.append((kind, symbol))
        if self.offline:
            raise PriceUnavailableError("Prices could not be fetched (offline).")
        if symbol not in self.listings:
            raise SymbolNotFoundError(f"Yahoo has no data for {symbol}.")
        return self.listings[symbol]

    def quote(self, symbol: str) -> Quote:
        listing = self._listing("quote", symbol)
        if listing.price is None:
            raise PriceUnavailableError(f"No current price for {symbol}.")
        return Quote(listing.info, listing.price, listing.info.currency or "", listing.as_of)

    def history(self, symbol: str, start: date, end: date) -> SymbolHistory:
        listing = self._listing("history", symbol)
        ranges = tuple(
            DailyRange(day, low, high) for day, (low, high) in sorted(listing.ranges.items()) if start <= day <= end
        )
        if not ranges:
            # Like Yahoo: a listing without prices in the period looks the same as a missing one.
            raise SymbolNotFoundError(f"Yahoo has no data for {symbol}.")
        return SymbolHistory(listing.info, ranges)

    def search(self, query: str) -> Sequence[SymbolInfo]:
        self.calls.append(("search", query))
        if self.offline:
            raise PriceUnavailableError("Search could not be reached (offline).")
        return self.search_results.get(query, ())
