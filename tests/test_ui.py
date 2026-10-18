"""The main window and its dialogs, driven as a user drives them, without a screen and with a mocked provider.

The tests act as the user does: they drop files on the window, click buttons, double-click cells,
type and press keys. They check only what the user can see (texts, pages, marks, the files on
disk), never the window's inner workings. What stands in for the user's own answers is replaced:
the file dialogs, "Confirm your investment" (``ask_listing``), "Enter its value" (``ask_value``),
showing Settings (``show_settings``, so that the test acts inside it) and the error boxes. Rule
numbers refer to tests/SPECIFICATION.md. No network.
"""

import gc
import os
import time
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
QtWidgets = pytest.importorskip("PySide6.QtWidgets", exc_type=ImportError)

from fakes import SAMPLE, FakeProvider, ranges_around, write_statement  # noqa: E402
from PySide6.QtCore import QEvent, QMimeData, QPoint, QPointF, Qt, QTimer, QUrl  # noqa: E402
from PySide6.QtGui import QDragEnterEvent, QDropEvent  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from statement_tracker import __version__  # noqa: E402
from statement_tracker.core.calculator import buys_by_ticker  # noqa: E402
from statement_tracker.core.errors import UserFacingError  # noqa: E402
from statement_tracker.core.models import MarketPrice  # noqa: E402
from statement_tracker.core.parser import load_statement  # noqa: E402
from statement_tracker.core.ticker_resolver import resolve_ticker  # noqa: E402
from statement_tracker.price_cache import PriceCache  # noqa: E402
from statement_tracker.settings import Settings, TickerChoice  # noqa: E402
from statement_tracker.ui import theme  # noqa: E402
from statement_tracker.ui.dialogs import ChooseListingDialog, EnterValueDialog, ValueChoice  # noqa: E402
from statement_tracker.ui.main_window import MainWindow  # noqa: E402
from statement_tracker.ui.settings_dialog import SettingsDialog  # noqa: E402

EQQB_BUYS = buys_by_ticker(load_statement(SAMPLE))["EQQB"]
TODAY = date(2026, 10, 12)
YOUR_VALUE = "(your value)"
"""How the window marks a value the user entered."""
LAST_KNOWN = " *"
"""How the window marks a value at the last price found on an earlier day."""


# -- the app and its windows --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def app():
    return QtWidgets.QApplication.instance() or QtWidgets.QApplication([])


@pytest.fixture(autouse=True)
def close_windows_afterwards(app):
    """Destroy each test's windows and dialogs when it ends.

    Left to Python's garbage collector, a window from one test can be destroyed at a random
    moment in the middle of a later test's event handling, which crashes Qt.
    """
    yield
    for widget in QtWidgets.QApplication.topLevelWidgets():
        widget.close()
        widget.deleteLater()
    QtWidgets.QApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete.value)
    gc.collect()


def provider_with(*symbols: str) -> FakeProvider:
    provider = FakeProvider()
    for symbol in symbols:
        provider.add(symbol, "EUR", price="400", ranges=ranges_around(EQQB_BUYS), name="Nasdaq fund", exchange="XETRA")
    return provider


def make_window(
    tmp_path: Path, provider: FakeProvider, now: datetime | None = None, today: date = TODAY
) -> tuple[MainWindow, list[UserFacingError]]:
    """The app's window on its own settings in ``tmp_path``, as it opens when the app starts."""
    extra = {"now": lambda: now} if now else {}
    window = MainWindow(Settings.load(tmp_path / "settings.json"), lambda: provider, lambda: today, **extra)
    errors: list[UserFacingError] = []
    window.show_error = errors.append
    window.ask_retry = lambda error: bool(errors.append(error))  # record it and answer "Not now"
    return window, errors


def settle(app) -> None:
    """Let work queued for "once the current event is handled" run."""
    for _ in range(3):
        app.processEvents()


def wait_until(app, condition, seconds: float = 20) -> None:
    deadline = time.monotonic() + seconds
    while not condition() and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    app.processEvents()
    assert condition(), "waited in vain"


def running(window: MainWindow) -> bool:
    """Whether the status line says the app is at work (shown in the theme's busy colour)."""
    return window.status_label.property("state") == "busy"


def run(app, window: MainWindow) -> None:
    """Press the green step 3 button and wait until the run is over."""
    assert window.go_button.isEnabled()
    window.go_button.click()
    wait_until(app, lambda: not running(window))


def drop(window: QtWidgets.QWidget, *paths: Path) -> None:
    """Drag files from the file manager and drop them on the window."""
    import warnings

    data = QMimeData()
    data.setUrls([QUrl.fromLocalFile(str(path)) for path in paths])
    where, action = window.rect().center(), Qt.DropAction.CopyAction
    buttons, keys = Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier
    with warnings.catch_warnings():  # PySide6 marks the only constructor it has as deprecated
        warnings.simplefilter("ignore", DeprecationWarning)
        entered = QDragEnterEvent(where, action, data, buttons, keys)
    QtWidgets.QApplication.sendEvent(window, entered)  # the file is dragged over the window...
    assert entered.isAccepted(), "the window does not take dropped files"
    QtWidgets.QApplication.sendEvent(window, QDropEvent(QPointF(where), action, data, buttons, keys))  # ...and let go


def button(widget: QtWidgets.QWidget, label: str) -> QtWidgets.QPushButton:
    """The button that reads ``label`` (without its keyboard shortcut's &)."""
    return next(found for found in widget.findChildren(QtWidgets.QPushButton) if found.text().replace("&", "") == label)


def cells(table: QtWidgets.QTableWidget) -> list[list[str]]:
    return [[table.item(row, column).text() for column in range(table.columnCount())] for row in range(table.rowCount())]


def cell_centre(table: QtWidgets.QTableWidget, row: int, column: int) -> QPoint:
    return table.visualItemRect(table.item(row, column)).center()


def bring_to_front(app, widget: QtWidgets.QWidget) -> None:
    """Show ``widget``'s window and make it the active one, as clicking into it does."""
    window = widget.window()
    window.show()
    window.raise_()
    window.activateWindow()
    QTest.qWaitForWindowActive(window, 1000)
    settle(app)


def double_click(app, table: QtWidgets.QTableWidget, row: int, column: int) -> None:
    bring_to_front(app, table)
    # A real double-click is a click and then a second one; QTest's double-click alone is only the second.
    where = cell_centre(table, row, column)
    QTest.mouseClick(table.viewport(), Qt.MouseButton.LeftButton, pos=where)
    QTest.mouseDClick(table.viewport(), Qt.MouseButton.LeftButton, pos=where)
    settle(app)


def select(app, table: QtWidgets.QTableWidget, row: int) -> None:
    bring_to_front(app, table)
    QTest.mouseClick(table.viewport(), Qt.MouseButton.LeftButton, pos=cell_centre(table, row, 0))
    settle(app)


def type_in_editor(app, table: QtWidgets.QTableWidget, text: str | None) -> None:
    """Type ``text`` over what the open cell editor shows and press Enter, or press Esc for None."""
    editor = next(editor for editor in table.findChildren(QtWidgets.QLineEdit) if editor.isVisible())
    if text is None:
        QTest.keyClick(editor, Qt.Key.Key_Escape)
    else:
        editor.selectAll()
        QTest.keyClick(editor, Qt.Key.Key_Delete)
        QTest.keyClicks(editor, text)
        QTest.keyClick(editor, Qt.Key.Key_Return)
    settle(app)


def editor_text(table: QtWidgets.QTableWidget) -> str:
    return next(editor for editor in table.findChildren(QtWidgets.QLineEdit) if editor.isVisible()).text()


def press(app, window: QtWidgets.QWidget, key: Qt.Key, modifiers=Qt.KeyboardModifier.NoModifier) -> None:
    """Press a key in ``window``, the active window, where its focus is."""
    bring_to_front(app, window)
    QTest.keyClick(QtWidgets.QApplication.focusWidget() or window, key, modifiers)
    settle(app)


def undo(app, window: QtWidgets.QWidget) -> None:
    press(app, window, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)


def redo(app, window: QtWidgets.QWidget) -> None:
    press(app, window, Qt.Key.Key_Y, Qt.KeyboardModifier.ControlModifier)


def in_settings(app, window: MainWindow, act) -> None:
    """Press Settings, do ``act(dialog)`` in it as the user would, and close it."""

    def shown(dialog: SettingsDialog) -> None:
        dialog.show()
        settle(app)
        act(dialog)
        dialog.close()

    window.show_settings = shown
    button(window, "Settings").click()
    settle(app)


def when_a_box_opens(act) -> list:
    """Have ``act(box)`` done to the next message box the app opens, as the user would answer it,
    and return the list that receives its text. The box's own buttons are clicked: nothing is replaced."""
    seen: list = []
    timer = QTimer()

    def check() -> None:
        box = QtWidgets.QApplication.activeModalWidget()
        if isinstance(box, QtWidgets.QMessageBox) and box.isVisible():
            timer.stop()
            seen.append("\n".join((box.text(), box.informativeText())))
            try:
                act(box)
            except Exception as problem:  # e.g. no such button: close the box, so the test fails instead of hanging
                seen.append(problem)
                box.reject()

    timer.timeout.connect(check)
    timer.start(10)
    seen.append(timer)  # kept alive until the test reads the text
    return seen


def text_of(seen: list) -> str:
    """The text of the box answered, after raising whatever went wrong answering it."""
    for item in seen:
        if isinstance(item, Exception):
            raise item
    assert any(isinstance(item, str) for item in seen), "no message box opened"
    return "\n".join(item for item in seen if isinstance(item, str))


def click(box: QtWidgets.QMessageBox, label: str) -> None:
    next(found for found in box.buttons() if found.text().replace("&", "") == label).click()


def shows(label: QtWidgets.QLabel, path: Path) -> bool:
    """True if the label shows the file's name and its folder, with the full path as tooltip."""
    return path.name in label.text() and str(path.parent) in label.text() and label.toolTip() == str(path)


def pixel(widget: QtWidgets.QWidget, x: int, y: int):
    return widget.grab().toImage().pixelColor(x, y)


def amount(text: str, number: str) -> bool:
    """Whether a shown amount starts with ``number``, written with a full stop or a comma."""
    return text.startswith((number, number.replace(".", ",")))


# -- F1, F2: the three steps -------------------------------------------------------------------------


def test_step_3_waits_for_a_statement(app, tmp_path: Path) -> None:
    # Spec F1.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))

    assert not window.go_button.isEnabled()
    assert window.no_spreadsheet_option.isChecked()
    assert window.go_button.text() == "Step 3: show my investments"

    drop(window, SAMPLE)

    assert window.go_button.isEnabled()
    assert shows(window.statement_label, SAMPLE)


def test_step_2_shows_the_choice_as_options_and_the_files_as_actions(
    app, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Spec F1, F2.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)
    existing, new = button(window, "Choose my spreadsheet..."), button(window, "Create a new one...")

    def state() -> tuple[bool, bool, bool]:
        """(no spreadsheet chosen, spreadsheet chosen, step 3 possible)"""
        return window.no_spreadsheet_option.isChecked(), window.spreadsheet_option.isChecked(), window.go_button.isEnabled()

    def files_shown() -> bool:
        return existing.isVisibleTo(window) and window.spreadsheet_label.isVisibleTo(window)

    # The file buttons are actions: never checkable, so never highlighted as if they were the choice.
    assert not existing.isCheckable() and not new.isCheckable()
    # Without a spreadsheet they mean nothing, so they are not shown.
    assert state() == (True, False, True) and not files_shown()

    # Choosing "Update my Excel spreadsheet" shows them, and asks for a file before step 3 can run.
    window.spreadsheet_option.click()
    assert state() == (False, True, False) and files_shown()
    assert window.spreadsheet_badge.text().startswith("To do")

    # A cancelled file dialog changes nothing.
    monkeypatch.setattr(QtWidgets.QFileDialog, "getOpenFileName", lambda *args, **kwargs: ("", ""))
    existing.click()
    assert state() == (False, True, False)

    new_file = tmp_path / "new.xlsx"
    monkeypatch.setattr(QtWidgets.QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(new_file), ""))
    new.click()
    assert state() == (False, True, True) and "will be created" in window.spreadsheet_label.text()

    # Back to no spreadsheet and then to the spreadsheet again: the file is still there.
    window.no_spreadsheet_option.click()
    assert state() == (True, False, True) and not files_shown()
    window.spreadsheet_option.click()
    assert state() == (False, True, True) and new_file.name in window.spreadsheet_label.text()

    # Dropping an Excel file on the window switches to the spreadsheet too.
    window.no_spreadsheet_option.click()
    drop(window, new_file)
    assert state() == (False, True, True) and files_shown()


def test_each_step_says_whether_it_is_done_optional_or_remembered(app, tmp_path: Path) -> None:
    # Spec F1.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))

    assert window.statement_badge.text().startswith("To do")
    assert window.spreadsheet_badge.text() == "Optional. You can skip this step."

    drop(window, SAMPLE, tmp_path / "mine.xlsx")
    assert window.statement_badge.text() == "✓ Done"
    assert window.spreadsheet_badge.text() == "✓ Done"

    # Next time the app starts, step 2 is already taken care of and says so, with nothing to press.
    reopened, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    assert reopened.spreadsheet_option.isChecked() and not reopened.no_spreadsheet_option.isChecked()
    assert reopened.spreadsheet_badge.text() == "✓ Already done: remembered from last time."
    assert shows(reopened.spreadsheet_label, tmp_path / "mine.xlsx")

    reopened.no_spreadsheet_option.click()
    assert reopened.spreadsheet_badge.text() == "Optional. You can skip this step."
    assert Settings.load(tmp_path / "settings.json").spreadsheet_path is None


def test_dropped_files_go_to_the_right_step_and_the_folder_is_remembered(app, tmp_path: Path) -> None:
    # Spec F2.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))

    drop(window, tmp_path / "book.xlsx")
    drop(window, SAMPLE)

    assert shows(window.spreadsheet_label, tmp_path / "book.xlsx")
    assert shows(window.statement_label, SAMPLE)
    assert Settings.load(tmp_path / "settings.json").last_folder == SAMPLE.parent


def test_file_dialogs_open_in_the_start_folder_until_a_folder_is_remembered(
    app, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Spec F2.
    project = tmp_path / "project"
    project.mkdir()
    monkeypatch.chdir(project)
    opened_in: list[str] = []

    def file_dialog(parent, caption, folder, *rest, **kwargs):
        opened_in.append(folder)
        return "", ""

    monkeypatch.setattr(QtWidgets.QFileDialog, "getOpenFileName", file_dialog)
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    choose = button(window, "Choose statement file...")

    choose.click()
    assert opened_in[-1] == str(project)  # run from source: where the app was started

    drop(window, SAMPLE)
    choose.click()
    assert opened_in[-1] == str(SAMPLE.parent)  # afterwards: the last folder used

    # The installed app, on a computer where nothing is remembered yet: somewhere of the user's, not the app's.
    monkeypatch.setattr("sys.frozen", True, raising=False)
    (tmp_path / "fresh").mkdir()
    installed, _ = make_window(tmp_path / "fresh", provider_with("EQQB.DE"))
    button(installed, "Choose statement file...").click()
    assert opened_in[-1] != str(project) and Path(opened_in[-1]).exists()


def test_help_page_gives_the_revolut_steps_and_asks_for_the_whole_history(app, tmp_path: Path) -> None:
    # Spec F11.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    box = when_a_box_opens(lambda box: click(box, "OK"))

    button(window, "Where do I get this file?").click()

    shown = text_of(box)
    for step in ("Invest", "More", "Documents", "Brokerage account", "Account statement", "CSV", "All time"):
        assert step in shown
    assert "whole history" in shown
    assert "PLACEHOLDER" not in shown


def test_about_shows_the_version_privacy_and_where_new_versions_are(app, tmp_path: Path) -> None:
    # Spec F11.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    box = when_a_box_opens(lambda box: click(box, "OK"))

    button(window, "About and privacy").click()

    shown = text_of(box)
    assert f"version {__version__}" in shown
    assert "https://github.com/TiDeane/Statement-Portfolio-Tracker/releases" in shown
    assert "Yahoo" in shown  # what leaves the computer


# -- F3, F4: the result -----------------------------------------------------------------------------


def test_the_result_has_its_own_page_in_the_same_window(app, tmp_path: Path) -> None:
    # Spec F3.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    window.show()
    settle(app)
    assert window.height() < 600  # just the three steps, no reserved empty space
    assert window.steps_page.isVisible() and not window.result_page.isVisible()
    drop(window, SAMPLE)
    settle(app)
    start_height = window.height()

    run(app, window)
    settle(app)

    assert window.result_page.isVisible() and not window.steps_page.isVisible()
    assert window.table.isVisible() and window.table.rowCount() == 2
    assert window.height() == start_height  # the same window, not a longer one

    window.back_button.click()
    assert window.steps_page.isVisible()
    window.go_button.click()  # step 3 again: back to the result page
    assert window.result_page.isVisible()
    assert window.back_button.isEnabled() and "Stop" in window.back_button.toolTip()  # it stops a run
    wait_until(app, lambda: not running(window))
    assert window.back_button.isEnabled() and "Stop" not in window.back_button.toolTip()

    # A file dropped on the result page is taken by the steps, which are shown again.
    drop(window, tmp_path / "mine.xlsx")
    assert window.steps_page.isVisible()


def test_the_window_fits_the_steps_until_there_is_a_result(app, tmp_path: Path) -> None:
    # Spec F10.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    window.show()
    settle(app)
    without_files = window.height()

    window.spreadsheet_option.click()
    settle(app)
    with_files = window.height()
    window.no_spreadsheet_option.click()
    settle(app)

    assert with_files > without_files  # grown for the file buttons...
    assert window.height() == without_files  # ...and back, with no empty band left at the bottom

    drop(window, SAMPLE)
    run(app, window)
    with_result = window.height()
    window.spreadsheet_option.click()
    window.no_spreadsheet_option.click()
    settle(app)
    assert window.height() >= with_result  # with a result shown, the window is never shrunk


def test_statement_alone_shows_the_investments(app, tmp_path: Path) -> None:
    # Spec F4.
    window, errors = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)

    run(app, window)

    assert errors == []
    headings = [window.table.horizontalHeaderItem(column).text() for column in range(window.table.columnCount())]
    assert headings == ["Investment", "Invested", "% of invested", "Worth now", "% of worth now", "Gain or loss"]
    eqqb, total = cells(window.table)
    assert eqqb[0] == "(EQQB) Nasdaq fund"
    assert "541" in eqqb[1] and "EUR" in eqqb[1]  # the statement's currency
    assert "589" in eqqb[3] and eqqb[3][0].isdigit()  # 1.47328346 shares at 400; the value itself has no arrow
    assert eqqb[5].startswith("▲ +47") and "(+8" in eqqb[5]  # 589.31 against 541.49
    assert eqqb[2].startswith("100") and eqqb[4].startswith("100")  # both shares: the only holding
    assert total[0] == "Total"
    assert window.status_label.text() == "Done. Your investments are shown below."
    assert list(tmp_path.glob("*.xlsx")) == []


@pytest.mark.parametrize(
    ("language", "country", "invested", "gain", "month"),
    [
        ("English", "UnitedStates", "12,345.60 EUR", ("+654.90", "5.30 %"), "October"),
        ("Portuguese", "Portugal", "345,60 EUR", ("+654,90", "5,30 %"), "outubro"),  # 12 345,60: spaced thousands
    ],
)
def test_numbers_and_dates_are_shown_as_the_computer_is_set_up(
    app, tmp_path: Path, language: str, country: str, invested: str, gain: tuple[str, str], month: str
) -> None:
    # Spec F4: the words stay English; numbers and dates follow the computer's own settings.
    from PySide6.QtCore import QLocale

    before = QLocale()
    QLocale.setDefault(QLocale(getattr(QLocale.Language, language), getattr(QLocale.Country, country)))
    try:
        noon = datetime(TODAY.year, TODAY.month, TODAY.day, 12).astimezone()
        window, _ = make_window(tmp_path, FakeProvider(), now=noon)  # Yahoo knows nothing: a value is entered
        window.ask_listing = lambda resolution: None
        drop(window, write_statement(
            tmp_path,
            "2026-01-05T09:00:00.000Z,,CASH TOP-UP,,,\"EUR 20,000\",EUR,1",
            "2026-01-05T10:00:00.000Z,ABCD,BUY - MARKET,100,EUR 123.45,\"EUR 12,345.60\",EUR,1",
        ))  # fmt: skip
        run(app, window)
        window.ask_value = lambda choices, selected, currency: ("ABCD", Decimal("13000.50"))
        button(window, "Enter its value...").click()

        abcd = cells(window.table)[0]
        assert abcd[1].endswith(invested) and "12,345.60" not in abcd[1].replace(invested, "")
        assert all(part in abcd[5] for part in gain)  # 13,000.50 against 12,345.60
        assert month in window.table.item(0, 3).toolTip()  # "Value you entered on <the day>"
    finally:
        QLocale.setDefault(before)


def test_gains_and_losses_have_their_own_sign_arrow_and_colour(app, tmp_path: Path) -> None:
    # Spec F4: not by colour alone.
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120")
    provider.add("WXYZ.DE", "EUR", price="40")
    window, _ = make_window(tmp_path, provider)
    window.ask_listing = lambda resolution: f"{resolution.ticker}.DE"
    drop(window, write_statement(
        tmp_path,
        "2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,EUR 2000,EUR,1",
        "2025-01-06T10:00:00.000Z,ABCD,BUY - MARKET,10,EUR 100.00,EUR 1000,EUR,1",
        "2025-01-20T10:00:00.000Z,WXYZ,BUY - MARKET,4,EUR 50.00,EUR 200,EUR,1",
    ))  # fmt: skip

    run(app, window)

    gain, loss = window.table.item(0, 5), window.table.item(1, 5)
    assert gain.text().startswith("▲ +200") and loss.text().startswith("▼ ")
    assert "40" in loss.text() and "+" not in loss.text()
    gain_colour, loss_colour = gain.foreground().color(), loss.foreground().color()
    assert gain_colour.green() > gain_colour.red() and loss_colour.red() > loss_colour.green()


def test_light_and_dark_themes_and_a_green_step_3_button(app, tmp_path: Path) -> None:
    # Spec F10.
    try:
        window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
        drop(window, SAMPLE)
        window.show()
        for dark in (False, True):
            theme.apply(app, dark=dark)
            settle(app)
            background = pixel(window, 2, window.height() - 2)
            assert (background.lightness() < 80) if dark else (background.lightness() > 180)
            go = window.go_button
            green = pixel(go, 4, go.height() // 2)
            assert green.green() > green.red() + 40 and green.green() > green.blue()
    finally:
        app.setStyleSheet("")


def test_many_investments_scroll_in_the_table_without_lengthening_the_window(app, tmp_path: Path) -> None:
    # Spec F10.
    tickers = [f"T{number:02d}" for number in range(20)]
    provider = FakeProvider()
    for ticker in tickers:
        provider.add(f"{ticker}.DE", "EUR", price="12")
    lines = ["2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,EUR 5000,EUR,1"] + [
        f"2025-01-06T10:00:{number:02d}.000Z,{ticker},BUY - MARKET,1,EUR 10.00,EUR 10,EUR,1"
        for number, ticker in enumerate(tickers)
    ]
    window, _ = make_window(tmp_path, provider)
    window.ask_listing = lambda resolution: f"{resolution.ticker}.DE"
    window.show()
    drop(window, write_statement(tmp_path, *lines))
    settle(app)
    start_height = window.height()

    run(app, window)
    settle(app)

    table = window.table
    assert table.rowCount() == 21  # twenty investments and the total
    assert window.height() == start_height
    assert table.height() < table.horizontalHeader().height() + table.rowCount() * table.rowHeight(0)
    assert table.verticalScrollBar().maximum() > 0  # the rest is reached by scrolling


def test_a_smaller_window_shrinks_the_table_instead_of_covering_it(app, tmp_path: Path) -> None:
    # Spec F10.
    window, _ = make_window(tmp_path, FakeProvider(offline=True))  # no price: the line with buttons shows
    window.show()
    drop(window, SAMPLE)
    run(app, window)
    settle(app)
    table = window.table
    header, rows = table.horizontalHeader().height(), [table.rowHeight(row) for row in range(table.rowCount())]
    assert table.viewport().height() + header >= header + sum(rows)  # every row shows
    settings = button(window, "Settings")

    def top(widget) -> int:
        return widget.mapTo(window, widget.rect().topLeft()).y()

    def bottom(widget) -> int:
        return widget.mapTo(window, widget.rect().bottomLeft()).y()

    def table_keeps_two_rows() -> bool:
        return table.height() >= header + 2 * rows[0]

    window.resize(window.width(), 100)  # far smaller than everything needs
    settle(app)

    assert window.height() > 100
    assert bottom(table) < top(window.price_bar) < top(settings)  # nothing covers the table
    assert table_keeps_two_rows()

    # Narrow as well: wrapped text takes more lines, and the window still keeps room for everything.
    window.resize(300, 100)
    settle(app)

    assert window.width() > 300
    assert bottom(window.status_label) < top(table)
    assert bottom(table) < top(window.price_bar) < top(settings)
    assert bottom(settings) <= window.height()
    assert table_keeps_two_rows()

    # Some window managers ignore the minimum size (Weston's, as under WSL). The window then grows back to it.
    minimum = window.minimumSize()
    window.setMinimumSize(0, 0)  # what Qt enforces: nothing, as such a window manager...
    window.minimumSize = lambda: minimum  # ...while the window still knows its own minimum
    window.resize(300, 200)
    settle(app)

    assert window.size() == minimum
    del window.minimumSize

    # And if it cannot grow back either, the contents scroll rather than being drawn over each other.
    window.resize(300, 200)
    settle(app)

    assert window.height() == 200
    assert window.centralWidget().verticalScrollBar().isVisible()
    assert bottom(window.status_label) < top(table)
    assert bottom(table) < top(window.price_bar) < top(settings)
    assert table_keeps_two_rows()


def test_buttons_are_only_highlighted_when_reached_with_the_keyboard(
    app, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Spec F10: a clicked button does not keep the highlight; one reached with Tab has it.
    monkeypatch.setattr(QtWidgets.QFileDialog, "getOpenFileName", lambda *args, **kwargs: ("", ""))
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    window.show()
    window.activateWindow()
    settle(app)
    buttons = window.findChildren(QtWidgets.QPushButton)
    assert not any(found.hasFocus() for found in buttons)  # none when the window opens

    choose = button(window, "Choose statement file...")
    QTest.mouseClick(choose, Qt.MouseButton.LeftButton)
    settle(app)
    assert not choose.hasFocus()

    for _ in range(len(buttons) + 5):
        press(app, window, Qt.Key.Key_Tab)
        if choose.hasFocus():
            break
    assert choose.hasFocus()


def shortcut_letters(window: QtWidgets.QWidget) -> list[str]:
    """The Alt shortcut letters of everything visible in ``window`` (the letter after "&", "&&" being a plain &)."""
    import re

    texts = [found.text() for found in window.findChildren(QtWidgets.QAbstractButton) if found.isVisibleTo(window)]
    texts += [found.text() for found in window.findChildren(QtWidgets.QLabel) if found.isVisibleTo(window) and found.buddy()]
    letters = []
    for text in texts:
        match = re.search(r"(?<!&)&([^&\s])", text.replace("&&", ""))
        if match:
            letters.append(match.group(1).lower())
    return letters


def reached_with_tab(app, window: QtWidgets.QWidget) -> set:
    """Everything that gets the focus when Tab is pressed over and over in ``window``."""
    bring_to_front(app, window)
    reached = set()
    for _ in range(60):
        QTest.keyClick(QtWidgets.QApplication.focusWidget() or window, Qt.Key.Key_Tab)
        settle(app)
        reached.add(QtWidgets.QApplication.focusWidget())
    return reached


def check_keyboard(app, window: QtWidgets.QWidget, where: str) -> None:
    letters = shortcut_letters(window)
    doubled = sorted({letter for letter in letters if letters.count(letter) > 1})
    assert not doubled, f"{where}: Alt+{', Alt+'.join(doubled)} is the shortcut of more than one thing"
    buttons = [found for found in window.findChildren(QtWidgets.QPushButton)
               if found.isVisibleTo(window) and found.isEnabled()]  # fmt: skip
    missed = [found.text().replace("&", "") for found in buttons if found not in reached_with_tab(app, window)]
    assert not missed, f"{where}: Tab never reaches {missed}"


def test_everything_can_be_reached_from_the_keyboard(app, tmp_path: Path) -> None:
    # Product rule "keyboard navigable": on every page and in every window, Tab reaches each button, and
    # no two things visible together share an Alt shortcut letter.
    window, _ = make_window(tmp_path, FakeProvider(offline=True))  # no price: the result's own buttons show too
    window.show()
    drop(window, SAMPLE, tmp_path / "mine.xlsx")  # with a spreadsheet, its buttons show too
    check_keyboard(app, window, "the steps")

    run(app, window)
    check_keyboard(app, window, "the result")

    def in_the_dialog(dialog: SettingsDialog) -> None:
        select(app, dialog.table, 0)  # the buttons for the selected investment are enabled
        check_keyboard(app, dialog, "Settings")

    in_settings(app, window, in_the_dialog)

    listing = ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", provider_with("EQQB.DE", "EQQB.MI")))
    listing.show()
    check_keyboard(app, listing, '"Confirm your investment"')
    value = EnterValueDialog([ValueChoice("AAAA", "(AAAA) First fund", Decimal("3"), None)], "AAAA", "EUR")
    value.show()
    QTest.keyClicks(value.value, "100")  # "Use this value" is enabled
    check_keyboard(app, value, '"Enter its value"')


def test_progress_bar_moves_forward_through_each_listing_price_and_the_spreadsheet(app, tmp_path: Path) -> None:
    # Spec F3.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    seen: list[tuple[int, int, str]] = []
    window.progressed.connect(lambda done, total, text: seen.append((done, total, text)))
    drop(window, SAMPLE, tmp_path / "mine.xlsx")
    assert window.progress_bar.isHidden()

    run(app, window)

    counted = [(done, total) for done, total, _ in seen if total]
    # Getting ready, finding EQQB's listing (a first run), its price, and the spreadsheet: four steps.
    assert {total for _, total in counted} == {4}
    assert [done for done, _ in counted] == sorted(done for done, _ in counted)  # never backwards
    texts = " ".join(text for _, _, text in seen)
    assert "listing of EQQB" in texts and "current prices" in texts and "spreadsheet" in texts
    assert window.progress_bar.isHidden()  # hidden again once finished
    assert window.progress_bar.format() == "%p%"


def test_while_a_listing_is_asked_the_status_says_the_app_waits_for_the_user(app, tmp_path: Path) -> None:
    # Spec F3.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE", "EQQB.MI"))  # two listings fit: it asks
    shown_while_asked: list[str] = []

    def answer(resolution):
        shown_while_asked.append(window.status_label.text())
        return "EQQB.DE"

    window.ask_listing = answer
    drop(window, SAMPLE)
    run(app, window)

    assert len(shown_while_asked) == 1
    assert shown_while_asked[0].startswith("Waiting for you") and "EQQB" in shown_while_asked[0]
    assert window.status_label.text() == "Done. Your investments are shown below."


# -- C3, C4: confirming an investment ---------------------------------------------------------------


def test_ambiguous_investment_asks_the_user_and_remembers_the_answer(app, tmp_path: Path) -> None:
    # Spec C3, C4.
    window, errors = make_window(tmp_path, provider_with("EQQB.DE", "EQQB.MI"))
    asked = []

    def choose(resolution):
        asked.append(resolution.ticker)
        return "EQQB.MI"

    window.ask_listing = choose
    drop(window, SAMPLE)

    run(app, window)
    run(app, window)

    assert errors == [] and asked == ["EQQB"]  # not asked again
    saved = Settings.load(tmp_path / "settings.json").ticker_choices["EQQB"]
    assert (saved.symbol, saved.confirmed) == ("EQQB.MI", True)


def test_skipped_investment_shows_as_unavailable(app, tmp_path: Path) -> None:
    # Spec C3, F5.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE", "EQQB.MI"))
    window.ask_listing = lambda resolution: None  # "Skip this investment"
    drop(window, SAMPLE)

    run(app, window)

    assert cells(window.table)[0][3] == "unavailable"
    assert cells(window.table)[0][4] == "unavailable"  # no share of worth without a worth


def test_listing_dialog_returns_the_selection_or_the_typed_symbol(app) -> None:
    # Spec C3.
    dialog = ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", provider_with("EQQB.DE", "EQQB.MI")))
    first, second = dialog.findChildren(QtWidgets.QRadioButton)
    typed = dialog.findChild(QtWidgets.QLineEdit)

    assert first.isChecked() and dialog.chosen_symbol() == "EQQB.DE"
    second.click()
    assert dialog.chosen_symbol() == "EQQB.MI"
    QTest.keyClicks(typed, " EQQB.SW ")
    assert dialog.chosen_symbol() == "EQQB.SW"


def test_always_use_the_first_listing_is_offered_whenever_there_is_a_listing(app) -> None:
    # Spec C3, C5.
    several = ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", provider_with("EQQB.DE", "EQQB.MI")))
    several.show()
    # The buttons, from left to right: use the selection, always use the first, skip.
    labels = ["Use this one", "Always choose the first option", "Skip this investment"]
    found = [button(several, label) for label in labels]
    assert [found_button.mapTo(several, QPoint(0, 0)).x() for found_button in found] == sorted(
        found_button.mapTo(several, QPoint(0, 0)).x() for found_button in found
    )

    # "Always choose the first option" uses the first listing now, even with another one selected.
    several.findChildren(QtWidgets.QRadioButton)[1].click()
    button(several, "Always choose the first option").click()
    assert several.result() == QtWidgets.QDialog.DialogCode.Accepted
    assert (several.chosen_symbol(), several.always_use_first()) == ("EQQB.DE", True)

    used = ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", provider_with("EQQB.DE", "EQQB.MI")))
    used.findChildren(QtWidgets.QRadioButton)[1].click()
    button(used, "Use this one").click()
    assert (used.chosen_symbol(), used.always_use_first()) == ("EQQB.MI", False)

    def offers_first(dialog: ChooseListingDialog) -> bool:
        return button(dialog, "Always choose the first option").isVisibleTo(dialog)

    # Choosing a matched listing again from Settings: the first one fits, so it is offered too.
    assert offers_first(ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", provider_with("EQQB.DE"))))

    # "The app could not tell for sure": listings found, none fitting the purchases. Offered as well.
    far_off = {day: (low + 100, high + 100) for day, (low, high) in ranges_around(EQQB_BUYS).items()}
    unmatched = FakeProvider()
    unmatched.add("EQQB.DE", "EUR", price="400", ranges=far_off)  # exists, but no purchase fits its prices
    assert offers_first(ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", unmatched)))

    # Nothing found at all: there is no first listing to take.
    assert not offers_first(ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", FakeProvider())))


def test_nothing_found_asks_for_a_symbol_or_to_skip(app) -> None:
    # Spec C3.
    dialog = ChooseListingDialog(resolve_ticker("EQQB", EQQB_BUYS, "EUR", FakeProvider()))
    use = button(dialog, "Use this symbol")
    typed = dialog.findChild(QtWidgets.QLineEdit)

    # Only typing a symbol gives something to use; until then, skipping is the way on.
    assert not use.isEnabled()
    QTest.keyClicks(typed, "  ")
    assert not use.isEnabled()
    typed.clear()
    QTest.keyClicks(typed, "EQQB.SW")
    assert use.isEnabled() and dialog.chosen_symbol() == "EQQB.SW"


def test_always_using_the_first_listing_is_remembered_and_shown_in_settings(
    app, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Spec C5, G4.
    from statement_tracker.ui import dialogs

    monkeypatch.setattr(dialogs, "ask_listing", lambda parent, resolution: ("EQQB.DE", True))
    window, _ = make_window(tmp_path, provider_with("EQQB.DE", "EQQB.MI"))
    drop(window, SAMPLE)
    run(app, window)

    assert window.table.item(0, 0).text() == "(EQQB) Nasdaq fund"
    settings_path = tmp_path / "settings.json"
    assert Settings.load(settings_path).always_use_first_listing
    assert Settings.load(settings_path).ticker_choices["EQQB"].symbol == "EQQB.DE"

    def turn_off(dialog: SettingsDialog) -> None:
        assert dialog.always_first.isChecked()
        dialog.always_first.click()

    in_settings(app, window, turn_off)
    assert not Settings.load(settings_path).always_use_first_listing


# -- F9: problems -------------------------------------------------------------------------------------


def test_wrong_file_is_explained_in_plain_words(app, tmp_path: Path) -> None:
    # Spec A6, F9.
    window, errors = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, write_statement(tmp_path, "2025-01-06,ABCD,1", header="Date,Symbol,Amount"))

    run(app, window)

    assert len(errors) == 1
    assert "does not look like a Revolut stock statement" in errors[0].what_happened
    assert window.status_label.text() == "That did not work. Nothing was changed."
    assert window.go_button.isEnabled()


def test_copy_details_puts_what_is_needed_to_report_a_problem_on_the_clipboard(app, tmp_path: Path) -> None:
    # Spec F9: an unexpected failure, in the real error box; the user presses "Copy details".
    def broken_provider():
        raise RuntimeError("boom")

    window = MainWindow(Settings.load(tmp_path / "settings.json"), broken_provider, lambda: TODAY)
    drop(window, SAMPLE)
    QtWidgets.QApplication.clipboard().clear()
    box = when_a_box_opens(lambda box: click(box, "Copy details"))

    run(app, window)

    shown = text_of(box)
    assert "Something unexpected went wrong." in shown and "Copy details" in shown
    assert "Traceback" not in shown and "boom" not in shown  # no technical details in the box itself
    copied = QtWidgets.QApplication.clipboard().text()
    assert __version__ in copied and "Something unexpected went wrong." in copied
    assert "RuntimeError: boom" in copied and "Traceback" in copied  # all a developer needs
    assert window.status_label.text() == "That did not work. Nothing was changed."


def test_locked_spreadsheet_still_shows_the_investments(app, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # Spec F9, H2.
    window, errors = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)
    run(app, window)  # first run settles the ticker choice, so only the spreadsheet is saved below
    drop(window, tmp_path / "mine.xlsx")

    def locked(source, target) -> None:
        raise PermissionError(13, "being used by another process")

    monkeypatch.setattr("os.replace", locked)
    run(app, window)

    assert "open in another program" in errors[0].what_happened
    assert cells(window.table)[0][0] == "(EQQB) Nasdaq fund"
    assert "spreadsheet was not updated" in window.status_label.text()


def test_try_again_saves_the_spreadsheet_without_fetching_prices_again(
    app, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Spec F9.
    provider = provider_with("EQQB.DE")
    window, _ = make_window(tmp_path, provider)
    drop(window, SAMPLE)
    run(app, window)  # settles the ticker choice
    spreadsheet = tmp_path / "mine.xlsx"
    drop(window, spreadsheet)

    real_replace, attempts, asked = os.replace, [], []

    def locked_twice(source, target) -> None:
        if Path(target) == spreadsheet:  # other files (settings, last prices) are saved normally
            attempts.append(target)
        if Path(target) == spreadsheet and len(attempts) <= 2:
            raise PermissionError(13, "being used by another process")
        real_replace(source, target)

    def try_again(error: UserFacingError) -> bool:
        asked.append(error.what_happened)
        return True

    monkeypatch.setattr("os.replace", locked_twice)
    window.ask_retry = try_again
    provider.calls.clear()
    run(app, window)

    assert len(asked) == 2 and all("open in another program" in text for text in asked)
    assert [call for call in provider.calls if call[0] == "quote"] == [("quote", "EQQB.DE")]  # once, not per attempt
    assert spreadsheet.exists()
    assert "Your spreadsheet was created" in window.status_label.text()


def test_back_to_the_steps_stops_a_run_without_saving_anything(app, tmp_path: Path) -> None:
    # Spec F8.
    import threading

    let_go = threading.Event()

    class Slow(FakeProvider):
        """Yahoo taking its time: each price waits until the test lets it go."""

        def quote(self, symbol: str):
            let_go.wait(10)
            return super().quote(symbol)

    provider = Slow()
    provider.add("EQQB.DE", "EUR", price="400", ranges=ranges_around(EQQB_BUYS), name="Nasdaq fund")
    window, errors = make_window(tmp_path, provider)
    spreadsheet = tmp_path / "mine.xlsx"
    window.show()
    drop(window, SAMPLE, spreadsheet)
    window.go_button.click()
    wait_until(app, lambda: "current prices" in window.status_label.text(), seconds=10)  # it waits for the price

    window.back_button.click()

    assert window.steps_page.isVisible()
    assert window.go_button.isEnabled()  # a new run can start at once
    let_go.set()  # the price comes after all
    deadline = time.monotonic() + 1
    while time.monotonic() < deadline:  # time for the stopped run to end, and for anything it might do
        app.processEvents()
        time.sleep(0.01)
    assert window.steps_page.isVisible() and errors == []  # nothing of it is shown
    assert not spreadsheet.exists()  # nor written
    assert Settings.load(tmp_path / "settings.json").ticker_choices == {}  # nor saved
    assert not (tmp_path / "last_prices.json").exists()

    run(app, window)  # step 3 again, to the end this time
    assert window.result_page.isVisible() and spreadsheet.exists()


# -- H: the spreadsheet, from the window ------------------------------------------------------------


def test_with_a_spreadsheet_it_is_created_and_remembered(app, tmp_path: Path) -> None:
    # Spec F1, H1.
    window, errors = make_window(tmp_path, provider_with("EQQB.DE"))
    spreadsheet = tmp_path / "mine.xlsx"
    drop(window, SAMPLE, spreadsheet)
    assert window.go_button.text() == "Step 3: show my investments and update my spreadsheet"
    assert "will be created" in window.spreadsheet_label.text()

    run(app, window)

    assert errors == []
    assert "Your spreadsheet was created" in window.status_label.text()
    assert load_workbook(spreadsheet).sheetnames == ["Current Holdings", "Portfolio History"]
    assert load_workbook(spreadsheet)["Current Holdings"]["A2"].value == "(EQQB) Nasdaq fund"

    reopened, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    assert shows(reopened.spreadsheet_label, spreadsheet)


def test_step_2_says_whether_a_backup_is_made_and_settings_change_it(app, tmp_path: Path) -> None:
    # Spec H5, G4.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    new_file, existing = tmp_path / "new.xlsx", tmp_path / "mine.xlsx"
    Workbook().save(existing)

    drop(window, new_file)
    assert not window.backup_label.isVisibleTo(window)  # nothing to back up yet
    drop(window, existing)
    assert window.backup_label.isVisibleTo(window)
    assert "backup copy" in window.backup_label.text() and "the 3 most recent are kept" in window.backup_label.text()
    assert window.backup_label.property("state") == "on"  # in grey

    def no_backups(dialog: SettingsDialog) -> None:
        assert dialog.backups.value() == 3
        dialog.backups.selectAll()
        QTest.keyClicks(dialog.backups, "0")
        assert dialog.backups.text() == "0"

    in_settings(app, window, no_backups)

    assert Settings.load(tmp_path / "settings.json").backups_kept == 0
    assert "No backup copy will be made" in window.backup_label.text()
    assert window.backup_label.property("state") == "off"  # in red
    drop(window, new_file)
    assert window.backup_label.isVisibleTo(window) and window.backup_label.property("state") == "off"

    def one_copy(dialog: SettingsDialog) -> None:
        dialog.backups.selectAll()
        QTest.keyClicks(dialog.backups, "1")

    in_settings(app, window, one_copy)
    drop(window, existing)
    assert "only the most recent is kept" in window.backup_label.text()


def test_with_backups_turned_off_no_copy_is_made(app, tmp_path: Path) -> None:
    # Spec H5.
    settings = Settings.load(tmp_path / "settings.json")
    settings.backups_kept = 0
    settings.save()
    window, errors = make_window(tmp_path, provider_with("EQQB.DE"))
    spreadsheet = tmp_path / "mine.xlsx"
    Workbook().save(spreadsheet)
    drop(window, SAMPLE, spreadsheet)

    run(app, window)

    assert errors == []
    assert "Your spreadsheet was updated" in window.status_label.text()
    assert sorted(file.name for file in tmp_path.glob("*.xlsx")) == ["mine.xlsx"]


# -- F6: renaming ---------------------------------------------------------------------------------------


def test_an_investment_is_renamed_by_typing_in_its_cell(app, tmp_path: Path) -> None:
    # Spec F6.
    window, errors = make_window(tmp_path, provider_with("EQQB.DE"))
    spreadsheet = tmp_path / "mine.xlsx"
    window.show()
    drop(window, SAMPLE, spreadsheet)
    run(app, window)
    asked = []
    window.ask_value = lambda *question: asked.append(question)
    table = window.table
    assert table.item(0, 0).text() == "(EQQB) Nasdaq fund"

    double_click(app, table, 0, 0)
    assert editor_text(table) == "Nasdaq fund"  # the name alone, without the ticker
    type_in_editor(app, table, " Tech fund ")

    assert table.item(0, 0).text() == "(EQQB) Tech fund"
    assert Settings.load(tmp_path / "settings.json").index_names == {"EQQB": "Tech fund"}
    assert load_workbook(spreadsheet)["Current Holdings"]["A2"].value == "(EQQB) Tech fund"
    assert asked == [] and errors == []  # no value was asked for, nothing went wrong

    double_click(app, table, 0, 0)
    type_in_editor(app, table, None)  # Esc cancels
    assert table.item(0, 0).text() == "(EQQB) Tech fund"

    double_click(app, table, 0, 0)
    type_in_editor(app, table, "")  # empty: back to the name from Yahoo
    assert table.item(0, 0).text() == "(EQQB) Nasdaq fund"
    assert Settings.load(tmp_path / "settings.json").index_names == {}

    # Nothing else in the table is typed into: the other cells open no editor.
    for row, column in ((0, 3), (1, 0)):
        double_click(app, table, row, column)
        assert not any(editor.isVisible() for editor in table.findChildren(QtWidgets.QLineEdit))


def test_an_investment_is_renamed_with_f2_too(app, tmp_path: Path) -> None:
    # Spec F6: select the investment, press F2, type, Enter.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    window.show()
    drop(window, SAMPLE)
    run(app, window)
    select(app, window.table, 0)

    press(app, window, Qt.Key.Key_F2)
    assert editor_text(window.table) == "Nasdaq fund"
    type_in_editor(app, window.table, "Tech fund")

    assert window.table.item(0, 0).text() == "(EQQB) Tech fund"
    assert Settings.load(tmp_path / "settings.json").index_names == {"EQQB": "Tech fund"}


# -- D3, E, F5: values that are not today's from Yahoo ----------------------------------------------


def test_last_known_price_is_marked_in_the_table_and_explained(app, tmp_path: Path) -> None:
    # Spec D3, F5.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)
    run(app, window)
    assert "*" not in cells(window.table)[0][3]

    offline, _ = make_window(tmp_path, FakeProvider(offline=True))
    drop(offline, SAMPLE)
    run(app, offline)

    worth, total = offline.table.item(0, 3), offline.table.item(1, 3)
    assert worth.text().endswith(LAST_KNOWN) and "Last known price, from" in worth.toolTip()
    assert total.text().endswith(LAST_KNOWN) and "not up to date" in total.toolTip()


def test_prices_line_appears_only_when_a_price_is_missing(app, tmp_path: Path) -> None:
    # Spec F5.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    window.show()
    drop(window, SAMPLE)
    run(app, window)
    assert window.price_bar.isHidden()
    asked = []
    window.ask_value = lambda *question: asked.append(question)
    double_click(app, window.table, 0, 3)
    assert asked == []  # a current price from Yahoo: nothing to enter

    (tmp_path / "other").mkdir()  # no last known price there
    offline, _ = make_window(tmp_path / "other", FakeProvider(offline=True))
    drop(offline, SAMPLE)
    run(app, offline)
    assert not offline.price_bar.isHidden()
    assert button(offline, "Try again").isVisibleTo(offline) and button(offline, "Enter its value...").isVisibleTo(offline)
    worth = offline.table.item(0, 3)
    assert worth.text() == "unavailable"
    assert "Double-click to enter its value" in worth.toolTip()
    # In red, as in the spreadsheet: missing, and a value can be entered there. Only that cell.
    red = worth.background().color()
    assert worth.background().style() != Qt.BrushStyle.NoBrush and red.red() > red.green() and red.red() > red.blue()
    assert offline.table.item(0, 1).background().style() == Qt.BrushStyle.NoBrush
    assert offline.table.item(1, 3).background().style() == Qt.BrushStyle.NoBrush  # nor the Total


def test_investments_without_a_price_are_left_out_of_the_totals_and_named(app, tmp_path: Path) -> None:
    # Spec B8, F5.
    statement = write_statement(
        tmp_path,
        "2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,EUR 2000,EUR,1",
        "2025-01-06T10:00:00.000Z,EQQB,BUY - MARKET,1,EUR 380.00,EUR 381,EUR,1",
        "2025-01-20T10:00:00.000Z,ZQXJ,BUY - MARKET,2,EUR 50.00,EUR 100,EUR,1",  # not on Yahoo
        "2025-01-21T10:00:00.000Z,QQQZ,BUY - MARKET,2,EUR 50.00,EUR 100,EUR,1",  # nor this
        "2025-01-22T10:00:00.000Z,ZZZQ,BUY - MARKET,2,EUR 50.00,EUR 100,EUR,1",  # nor this
    )
    settings = Settings.load(tmp_path / "settings.json")
    settings.remember("EQQB", TickerChoice("EQQB.DE", "Nasdaq fund", confirmed=True))
    settings.save()

    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    window.ask_listing = lambda resolution: None  # "Skip this investment"
    drop(window, statement)
    run(app, window)

    shown = {row[0]: row for row in cells(window.table)}
    eqqb, zqxj, total = shown["(EQQB) Nasdaq fund"], shown["ZQXJ"], shown["Total"]
    assert zqxj[2] == zqxj[4] == "not in total" and zqxj[3] == "unavailable"
    named = window.price_bar_label.text()
    assert named.startswith("Yahoo Finance has no current price for ") and named.endswith(".")
    names = named.removeprefix("Yahoo Finance has no current price for ").removesuffix(".")
    assert sorted(names.replace(" and ", ", ").split(", ")) == ["QQQZ", "ZQXJ", "ZZZQ"]
    assert ", " in names and names.count(" and ") == 1  # "A, B and C"
    zqxj_row = [row[0] for row in cells(window.table)].index("ZQXJ")
    assert "left out of the totals" in window.table.item(zqxj_row, 3).toolTip()
    assert eqqb[2] == eqqb[4] == total[2] == total[4]  # 100 %: EQQB is all the totals cover
    assert total[1] == eqqb[1] and total[3] == eqqb[3]


def test_entered_value_fills_in_the_table_and_the_spreadsheet_as_typed(app, tmp_path: Path) -> None:
    # Spec E1, E2, E3, E5.
    noon = datetime(TODAY.year, TODAY.month, TODAY.day, 12).astimezone()
    window, errors = make_window(tmp_path, FakeProvider(offline=True), now=noon)
    spreadsheet = tmp_path / "mine.xlsx"
    window.show()
    drop(window, SAMPLE, spreadsheet)
    run(app, window)
    asked = []

    def answer(choices, selected, currency):
        asked.append(([(choice.label, choice.shares) for choice in choices], selected, currency))
        return "EQQB", Decimal("600")  # what all 1.47 shares are worth, not one share's price

    window.ask_value = answer
    double_click(app, window.table, 1, 3)  # the Total row: nothing to enter there
    assert asked == []
    double_click(app, window.table, 0, 3)

    assert asked == [([("EQQB", Decimal("1.47328346"))], "EQQB", "EUR")]
    assert errors == []
    worth, total = window.table.item(0, 3), window.table.item(1, 3)
    assert amount(worth.text(), "600.00")  # the number typed is the number shown
    assert worth.text().endswith(f" {YOUR_VALUE}") and "Value you entered on" in worth.toolTip()
    assert total.text().endswith(f" {YOUR_VALUE}") and "entered yourself" in total.toolTip()
    assert not window.price_bar.isHidden()  # the value can still be changed, or fetched again
    assert load_workbook(spreadsheet)["Current Holdings"]["D2"].value == pytest.approx(600)
    assert load_workbook(spreadsheet)["Portfolio History"].max_row == 2  # typed in today: today's row is written

    window.ask_value = lambda choices, selected, currency: ("EQQB", Decimal("123"))  # changed again
    button(window, "Enter its value...").click()
    assert amount(window.table.item(0, 3).text(), "123.00")


def test_an_entered_value_follows_the_shares_bought_since(app, tmp_path: Path) -> None:
    # Spec E4: the value entered is kept per share, so buying more shares raises it.
    bought = "2025-01-06T10:00:00.000Z,ZQXJ,BUY - MARKET,4,EUR 50.00,EUR 200,EUR,1"
    top_up = "2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,EUR 1000,EUR,1"
    window, errors = make_window(tmp_path, FakeProvider())  # Yahoo does not know ZQXJ
    window.ask_listing = lambda resolution: None  # "Skip this investment"
    drop(window, write_statement(tmp_path, top_up, bought))
    run(app, window)
    window.ask_value = lambda choices, selected, currency: ("ZQXJ", Decimal("400"))  # what the user types
    button(window, "Enter its value...").click()
    assert amount(window.table.item(0, 3).text(), "400.00")

    more = "2025-03-03T10:00:00.000Z,ZQXJ,BUY - MARKET,2,EUR 60.00,EUR 120,EUR,1"
    drop(window, write_statement(tmp_path, top_up, bought, more))
    run(app, window)

    worth = window.table.item(0, 3).text()
    assert amount(worth, "600.00") and worth.endswith(YOUR_VALUE)  # 6 shares at 100
    assert errors == []


def test_a_value_typed_into_the_spreadsheet_is_used_on_the_next_run(app, tmp_path: Path) -> None:
    # Spec E1.
    window, errors = make_window(tmp_path, FakeProvider(offline=True))
    spreadsheet = tmp_path / "mine.xlsx"
    drop(window, SAMPLE, spreadsheet)
    run(app, window)
    assert window.table.item(0, 3).text() == "unavailable"
    workbook = load_workbook(spreadsheet)
    workbook["Current Holdings"]["D2"] = 600  # typed by the user over "unavailable"
    workbook.save(spreadsheet)

    run(app, window)

    assert errors == []
    assert window.table.item(0, 3).text().endswith(YOUR_VALUE) and "600" in window.table.item(0, 3).text()


def test_enter_value_dialog_starts_from_the_value_known_and_returns_the_one_typed(app) -> None:
    # Spec E1, E7.
    last_known = MarketPrice(
        Decimal("100"), "BBBB.DE", "EUR", as_of=datetime(2026, 10, 3, 17, 30, tzinfo=UTC), last_known=True
    )
    choices = [
        ValueChoice("AAAA", "(AAAA) First fund", Decimal("3"), None),
        ValueChoice("BBBB", "(BBBB) Second fund", Decimal("2.5"), last_known),
    ]
    dialog = EnterValueDialog(choices, "BBBB", "EUR")
    use = button(dialog, "Use this value")

    assert dialog.investment.isVisibleTo(dialog) and dialog.investment.currentText().startswith("(BBBB)")
    assert amount(dialog.value.text(), "250")  # 2.5 shares at 100: the whole value
    assert use.isEnabled()

    dialog.investment.setCurrentIndex(0)
    assert dialog.value.text() == "" and not use.isEnabled()
    QTest.keyClicks(dialog.value, "abc")
    assert not use.isEnabled() and dialog.entered() is None
    dialog.value.clear()
    QTest.keyClicks(dialog.value, "512,34")
    assert use.isEnabled()
    assert dialog.entered() == ("AAAA", Decimal("512.34"))


@pytest.mark.parametrize(
    ("text", "value"),
    [
        ("512.34", "512.34"), ("512,34", "512.34"), (" 500 ", "500"), ("1,234.56", "1234.56"),
        ("1.234,56", "1234.56"), ("1 234,5", "1234.5"), ("0", None), ("-5", None), ("", None),
        ("abc", None), ("1.2.3", None), ("12 EUR", None), ("NaN", None),
    ],
)  # fmt: skip
def test_typed_values_are_read_with_a_comma_or_a_full_stop(app, text: str, value: str | None) -> None:
    # Spec E7: typed into "Enter its value", as the user types it.
    dialog = EnterValueDialog([ValueChoice("AAAA", "(AAAA) First fund", Decimal("3"), None)], "AAAA", "EUR")
    QTest.keyClicks(dialog.value, text)

    assert dialog.entered() == (("AAAA", Decimal(value)) if value else None)
    assert button(dialog, "Use this value").isEnabled() is (value is not None)


# -- G: Settings ----------------------------------------------------------------------------------------


def settings_row(dialog: SettingsDialog, ticker: str) -> list[str]:
    (row,) = [row for row in range(dialog.table.rowCount()) if dialog.table.item(row, 0).text() == ticker]
    return [dialog.table.item(row, column).text() for column in range(dialog.table.columnCount())]


def forget(app, dialog: SettingsDialog, label: str = "Forget", answer: str = "Yes") -> str:
    """Press Forget (or Forget all...) and answer the question; return the question."""
    box = when_a_box_opens(lambda box: click(box, answer))
    button(dialog, label).click()
    settle(app)
    return text_of(box)


def test_settings_rename_choose_again_and_forget_an_investment(app, tmp_path: Path) -> None:
    # Spec G1, G2, G5, G6.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)
    run(app, window)
    saved = tmp_path / "settings.json"

    def act(dialog: SettingsDialog) -> None:
        headings = [dialog.table.horizontalHeaderItem(column).text() for column in range(dialog.table.columnCount())]
        assert headings == ["Investment", "Name shown", "Listing used", "Status"]
        assert settings_row(dialog, "EQQB") == ["EQQB", "Nasdaq fund", "EQQB.DE (XETRA, EUR)", "Matched automatically"]
        assert "from Yahoo Finance" in dialog.table.item(0, 1).toolTip()  # Yahoo's name, not the user's
        assert not button(dialog, "Rename").isEnabled()  # nothing selected yet

        select(app, dialog.table, 0)
        button(dialog, "Rename").click()
        type_in_editor(app, dialog.table, " Tech fund ")
        assert settings_row(dialog, "EQQB")[1] == "Tech fund"
        assert "from Yahoo Finance" not in dialog.table.item(0, 1).toolTip()
        assert Settings.load(saved).index_names == {"EQQB": "Tech fund"}

        button(dialog, "Choose the listing next time").click()
        assert settings_row(dialog, "EQQB")[3] == "Asks you next time"
        assert not button(dialog, "Choose the listing next time").isEnabled()
        assert Settings.load(saved).ask_again == {"EQQB"}

        question = forget(app, dialog, answer="Cancel")
        assert question.startswith("Forget EQQB?") and dialog.table.rowCount() == 1  # cancelled
        select(app, dialog.table, 0)
        forget(app, dialog)
        # The empty table stays, headings and all, with the reason inside it.
        assert dialog.table.rowCount() == 0 and dialog.table.isVisible()
        assert any("No investments are remembered yet" in label.text() and label.isVisible()
                   for label in dialog.findChildren(QtWidgets.QLabel))  # fmt: skip
        assert Settings.load(saved).known_tickers() == []

    in_settings(app, window, act)
    assert window.table.item(0, 0).text() == "EQQB"  # the result shows it at once: forgotten, with its name


def test_settings_only_the_name_is_edited_in_the_table(app, tmp_path: Path) -> None:
    # Spec G2, F6.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)
    run(app, window)
    saved = tmp_path / "settings.json"

    def act(dialog: SettingsDialog) -> None:
        for column in (0, 2, 3):
            double_click(app, dialog.table, 0, column)
            assert not any(editor.isVisible() for editor in dialog.table.findChildren(QtWidgets.QLineEdit))

        def rename(text: str | None) -> None:
            double_click(app, dialog.table, 0, 1)
            type_in_editor(app, dialog.table, text)

        rename("Nasdaq fund")  # Yahoo's name, unchanged: stays Yahoo's
        assert Settings.load(saved).index_names == {}
        rename("Mine")
        assert Settings.load(saved).index_names == {"EQQB": "Mine"}
        rename(None)  # Esc cancels
        assert Settings.load(saved).index_names == {"EQQB": "Mine"}
        rename("  ")  # empty: back to Yahoo's name
        assert Settings.load(saved).index_names == {}
        assert dialog.table.item(0, 1).text() == "Nasdaq fund"

    in_settings(app, window, act)


def test_settings_forget_all_keeps_the_rest(app, tmp_path: Path) -> None:
    # Spec G3.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE, tmp_path / "mine.xlsx")
    run(app, window)
    saved = tmp_path / "settings.json"

    def act(dialog: SettingsDialog) -> None:
        question = forget(app, dialog, "Forget all...", answer="Cancel")
        assert question.startswith("Forget your one investment?") and dialog.table.rowCount() == 1

        forget(app, dialog, "Forget all...")
        assert Settings.load(saved).known_tickers() == [] and not button(dialog, "Forget all...").isEnabled()

    in_settings(app, window, act)
    assert Settings.load(saved).spreadsheet_path == tmp_path / "mine.xlsx"
    assert PriceCache.beside(saved).last_known("EQQB.DE", "EUR") is not None  # Yahoo's prices stay


def test_settings_forget_the_values_entered_too_but_not_the_prices_found(app, tmp_path: Path) -> None:
    # Spec G1, G2, G3.
    provider = provider_with("EQQB.DE")
    window, _ = make_window(tmp_path, provider)
    drop(window, SAMPLE)
    run(app, window)  # EQQB matched, and its price found
    provider.offline = True
    run(app, window)
    window.ask_value = lambda choices, selected, currency: ("EQQB", Decimal("600"))
    button(window, "Enter its value...").click()
    statement = write_statement(tmp_path, "2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,EUR 1000,EUR,1",
                                "2025-01-20T10:00:00.000Z,ZQXJ,BUY - MARKET,2,EUR 50.00,EUR 100,EUR,1")  # fmt: skip
    window.ask_listing = lambda resolution: None
    drop(window, statement)
    run(app, window)  # ZQXJ: never matched to a listing; only its value is kept
    window.ask_value = lambda choices, selected, currency: ("ZQXJ", Decimal("110"))
    button(window, "Enter its value...").click()

    def act(dialog: SettingsDialog) -> None:
        assert settings_row(dialog, "ZQXJ") == ["ZQXJ", "", "", "Not matched yet; value entered by you"]
        assert settings_row(dialog, "EQQB")[3] == "Matched automatically; value entered by you"

        select(app, dialog.table, 1)  # ZQXJ, after EQQB
        forget(app, dialog)
        assert [dialog.table.item(row, 0).text() for row in range(dialog.table.rowCount())] == ["EQQB"]

        forget(app, dialog, "Forget all...")
        assert dialog.table.rowCount() == 0

    in_settings(app, window, act)
    kept = PriceCache.beside(tmp_path / "settings.json")
    assert kept.entered_tickers() == set()
    assert kept.last_known("EQQB.DE", "EUR") is not None  # the price found on Yahoo stays


def test_forgetting_in_settings_shows_at_once_in_the_result(app, tmp_path: Path) -> None:
    # Spec G5.
    provider = provider_with("EQQB.DE")
    window, errors = make_window(tmp_path, provider)
    drop(window, SAMPLE)
    run(app, window)
    provider.offline = True
    run(app, window)  # Yahoo unreachable: the last price found
    window.ask_value = lambda choices, selected, currency: ("EQQB", Decimal("450"))
    button(window, "Enter its value...").click()
    assert window.table.item(0, 3).text().endswith(YOUR_VALUE)

    in_settings(app, window, lambda dialog: forget(app, dialog, "Forget all..."))

    assert errors == []
    assert window.table.item(0, 3).text().endswith(LAST_KNOWN)  # back to the last price found, not the one entered
    assert Settings.load(tmp_path / "settings.json").known_tickers() == []  # and the listing stays forgotten


def test_settings_highlight_the_whole_row_under_the_mouse(app, tmp_path: Path) -> None:
    # Spec G1: a list of investments, so the whole row under the mouse is highlighted, not one cell.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)
    run(app, window)

    def act(dialog: SettingsDialog) -> None:
        table = dialog.table
        far = table.visualItemRect(table.item(0, 3))
        before = table.viewport().grab().toImage().pixelColor(far.right() - 3, far.center().y())
        QTest.mouseMove(table.viewport(), cell_centre(table, 0, 1))  # the mouse over the name
        table.entered.emit(table.model().index(0, 1))  # what Qt reports for it, offscreen too
        settle(app)
        hovered = table.viewport().grab().toImage().pixelColor(far.right() - 3, far.center().y())
        assert hovered != before  # the status cell, far from the mouse, is highlighted too
        table.viewportEvent(QEvent(QEvent.Type.Leave))  # the mouse leaves the table
        settle(app)
        assert table.viewport().grab().toImage().pixelColor(far.right() - 3, far.center().y()) == before

    in_settings(app, window, act)


def test_settings_cannot_be_opened_while_running(app, tmp_path: Path) -> None:
    # Spec G5.
    window, _ = make_window(tmp_path, provider_with("EQQB.DE"))
    drop(window, SAMPLE)
    window.go_button.click()
    assert not button(window, "Settings").isEnabled()
    wait_until(app, lambda: not running(window))
    assert button(window, "Settings").isEnabled()


# -- F7: undo ---------------------------------------------------------------------------------------------


def test_settings_changes_to_investments_are_undone_and_redone(app, tmp_path: Path) -> None:
    # Spec F7.
    provider = provider_with("EQQB.DE")
    window, _ = make_window(tmp_path, provider)
    drop(window, SAMPLE)
    run(app, window)
    provider.offline = True
    run(app, window)  # Yahoo unreachable: a value can be entered
    window.ask_value = lambda choices, selected, currency: ("EQQB", Decimal("600"))
    button(window, "Enter its value...").click()
    saved = tmp_path / "settings.json"

    def act(dialog: SettingsDialog) -> None:
        double_click(app, dialog.table, 0, 1)
        type_in_editor(app, dialog.table, "Tech fund")
        undo(app, dialog)
        assert Settings.load(saved).index_names == {} and dialog.table.item(0, 1).text() == "Nasdaq fund"
        assert dialog.message_label.text() == "Undone: the new name of EQQB."
        redo(app, dialog)
        assert Settings.load(saved).index_names == {"EQQB": "Tech fund"}

        select(app, dialog.table, 0)
        forget(app, dialog)
        undo(app, dialog)  # the listing, the name and the value entered all come back
        assert settings_row(dialog, "EQQB") == [
            "EQQB", "Tech fund", "EQQB.DE (XETRA, EUR)", "Matched automatically; value entered by you",
        ]  # fmt: skip

        forget(app, dialog, "Forget all...")
        undo(app, dialog)
        assert Settings.load(saved).known_tickers() == ["EQQB"]

        # The tick box and the backups are not part of it: undoing goes past them and leaves them as they are.
        dialog.always_first.click()
        dialog.backups.selectAll()
        QTest.keyClicks(dialog.backups, "5")
        select(app, dialog.table, 0)  # out of the box, whose own Ctrl+Z is for its text
        undo(app, dialog)  # the redo of the name
        settings = Settings.load(saved)
        assert settings.index_names == {} and settings.always_use_first_listing and settings.backups_kept == 5
        assert dialog.always_first.isChecked() and dialog.backups.value() == 5
        undo(app, dialog)  # the value entered in the result, before Settings was opened
        undo(app, dialog)  # nothing left to undo: nothing happens
        assert Settings.load(saved).index_names == {}
        assert "value entered by you" not in settings_row(dialog, "EQQB")[3]

    in_settings(app, window, act)


def test_names_and_values_changed_in_the_result_are_undone_there_and_in_settings(app, tmp_path: Path) -> None:
    # Spec F7.
    provider = provider_with("EQQB.DE")
    window, errors = make_window(tmp_path, provider)
    spreadsheet = tmp_path / "mine.xlsx"
    window.show()
    drop(window, SAMPLE, spreadsheet)
    run(app, window)
    provider.offline = True
    run(app, window)  # Yahoo unreachable: a value can be entered
    double_click(app, window.table, 0, 0)
    type_in_editor(app, window.table, "Tech fund")
    window.ask_value = lambda choices, selected, currency: ("EQQB", Decimal("450"))
    button(window, "Enter its value...").click()
    assert window.table.item(0, 3).text().endswith(YOUR_VALUE)

    undo(app, window)
    assert window.table.item(0, 3).text().endswith(LAST_KNOWN)  # back to the last price found
    assert load_workbook(spreadsheet)["Current Holdings"]["D2"].comment.text.count("Last known price") == 1
    assert load_workbook(spreadsheet)["Current Holdings"]["A2"].value == "(EQQB) Tech fund"  # written again

    # Settings carries on with the same history: the rename made in the result is undone there.
    def undo_and_redo_there(dialog: SettingsDialog) -> None:
        undo(app, dialog)
        assert dialog.table.item(0, 1).text() == "Nasdaq fund"
        assert dialog.message_label.text() == "Undone: the new name of EQQB."
        redo(app, dialog)
        redo(app, dialog)  # and the value entered in the result, redone
        assert "value entered by you" in settings_row(dialog, "EQQB")[3]

    in_settings(app, window, undo_and_redo_there)
    assert window.table.item(0, 0).text() == "(EQQB) Tech fund"
    assert window.table.item(0, 3).text().endswith(YOUR_VALUE)

    window.back_button.click()
    undo(app, window)  # only on the result page, where the change shows
    assert Settings.load(tmp_path / "settings.json").index_names == {"EQQB": "Tech fund"}
    assert errors == []


def test_after_a_run_names_and_values_are_still_undone_but_forgetting_is_not(app, tmp_path: Path) -> None:
    # Spec F7.
    provider = provider_with("EQQB.DE")
    window, errors = make_window(tmp_path, provider)
    window.show()
    drop(window, SAMPLE)
    run(app, window)
    provider.offline = True
    run(app, window)  # Yahoo unreachable: a value can be entered
    saved = tmp_path / "settings.json"
    double_click(app, window.table, 0, 0)
    type_in_editor(app, window.table, "Tech fund")
    window.ask_value = lambda choices, selected, currency: ("EQQB", Decimal("450"))
    button(window, "Enter its value...").click()
    run(app, window)

    undo(app, window)
    assert not window.table.item(0, 3).text().endswith(YOUR_VALUE)
    assert window.table.item(0, 0).text() == "(EQQB) Tech fund"
    undo(app, window)
    assert window.table.item(0, 0).text() == "(EQQB) Nasdaq fund"

    in_settings(app, window, lambda dialog: (select(app, dialog.table, 0), forget(app, dialog)))
    provider.offline = False
    run(app, window)  # matches the listing again
    matched = Settings.load(saved).ticker_choices["EQQB"]

    def undo_there(dialog: SettingsDialog) -> None:
        undo(app, dialog)  # forgetting cannot be undone after a run: nothing happens
        assert dialog.message_label.text() == ""

    in_settings(app, window, undo_there)
    undo(app, window)
    assert Settings.load(saved).ticker_choices["EQQB"] == matched
    assert errors == []


def test_undoing_a_rename_after_a_run_brings_back_only_the_name(app, tmp_path: Path) -> None:
    # Spec F7: each step puts back only what it changed, not the listing chosen since.
    window, errors = make_window(tmp_path, provider_with("EQQB.DE", "EQQB.MI"))
    window.ask_listing = lambda resolution: "EQQB.DE"
    window.show()
    drop(window, SAMPLE)
    run(app, window)
    double_click(app, window.table, 0, 0)
    type_in_editor(app, window.table, "Tech fund")
    in_settings(app, window, lambda dialog: (select(app, dialog.table, 0),
                                             button(dialog, "Choose the listing next time").click()))  # fmt: skip
    window.ask_listing = lambda resolution: "EQQB.MI"  # picked again on the next run
    run(app, window)

    undo(app, window)  # the rename: choosing the listing again is past undoing now

    saved = Settings.load(tmp_path / "settings.json")
    assert saved.index_names == {} and saved.ticker_choices["EQQB"].symbol == "EQQB.MI" and saved.ask_again == set()
    assert window.table.item(0, 0).text() == "(EQQB) Nasdaq fund"
    assert errors == []


# -- weeks of use ----------------------------------------------------------------------------------------------


def test_weeks_of_use_keep_one_history_row_per_valid_day_and_the_newest_backups(app, tmp_path: Path) -> None:
    # Spec D3, E4-E6, H4, H5, over time: the app is opened most days, each time with a statement exported
    # that day; prices move, Yahoo is unreachable on some days, and the user fills in a value both ways.
    import os

    from fakes import computed

    rows = [  # (date, row): each statement has the rows dated before the day it is exported
        (date(2026, 8, 3), "2026-08-03T09:00:00.000Z,,CASH TOP-UP,,,EUR 3000,EUR,1"),
        (date(2026, 8, 3), "2026-08-03T10:00:00.000Z,ABCD,BUY - MARKET,10,EUR 100.00,EUR 1001,EUR,1"),
        (date(2026, 8, 4), "2026-08-04T10:00:00.000Z,WXYZ,BUY - MARKET,4,EUR 50.00,EUR 201,EUR,1"),
        (date(2026, 9, 4), "2026-09-04T10:00:00.000Z,ZQXJ,BUY - MARKET,2,EUR 50.00,EUR 100,EUR,1"),
        (date(2026, 9, 10), "2026-09-10T10:00:00.000Z,ABCD,BUY - MARKET,2,EUR 130.00,EUR 260,EUR,1"),
        (date(2026, 9, 18), "2026-09-18T10:00:00.000Z,WXYZ,SELL - MARKET,2,EUR 55.00,EUR 110,EUR,1"),
    ]
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", ranges={date(2026, 8, 3): (Decimal(99), Decimal(101)),
                                           date(2026, 9, 10): (Decimal(129), Decimal(131))})  # fmt: skip
    provider.add("WXYZ.DE", "EUR", ranges={date(2026, 8, 4): (Decimal(49), Decimal(51))})
    spreadsheet = tmp_path / "mine.xlsx"
    history_as_shown: dict[date, list] = {}  # each day's history row, as Excel showed it at the end of that day
    tables: dict[date, list[list[str]]] = {}

    def open_the_app(day: date, offline: bool = False, *, before=None, after=None) -> None:
        number = (day - date(2026, 9, 1)).days
        provider.offline = offline
        for symbol, start in (("ABCD.DE", 120), ("WXYZ.DE", 60), ("ZQXJ.DE", 70)):
            if symbol in provider.listings:
                provider.listings[symbol].price = Decimal(start + number)
        if before:
            before()
        noon = datetime(day.year, day.month, day.day, 12).astimezone()
        window, errors = make_window(tmp_path, provider, now=noon, today=day)
        window.ask_listing = lambda resolution: None  # ZQXJ: nothing found, so "Skip this investment"
        statement = write_statement(tmp_path, *(row for when, row in rows if when < day))
        if not spreadsheet.exists():
            drop(window, statement, spreadsheet)  # the first day: step 2 too; it is remembered from then on
        else:
            drop(window, statement)
            assert shows(window.spreadsheet_label, spreadsheet)
        run(app, window)
        if after:
            after(window)
        assert errors == []
        tables[day] = cells(window.table)
        sheet = load_workbook(spreadsheet)["Portfolio History"]
        dates = [cell.value.date() for cell in sheet["A"][1:]]
        shown = computed(spreadsheet, "Portfolio History")[1:]
        # Earlier days' rows show what they showed on their day, whatever happened since.
        for when, row in zip(dates, shown):
            if when in history_as_shown and when != day:
                assert row == history_as_shown[when], f"the row of {when} changed on {day}"
        history_as_shown.update(zip(dates, shown))
        window.close()

    def worth(day: date, ticker: str) -> str:
        return next(row[3] for row in tables[day] if row[0] == ticker or row[0].startswith(f"({ticker})"))

    def enter_value(window: MainWindow) -> None:
        window.ask_value = lambda choices, selected, currency: ("ZQXJ", Decimal("120"))  # all 2 shares
        button(window, "Enter its value...").click()

    def type_into_the_sheet() -> None:
        workbook = load_workbook(spreadsheet)
        sheet = workbook["Current Holdings"]
        (row,) = [row for row in range(2, sheet.max_row + 1) if str(sheet.cell(row=row, column=1).value).endswith("ZQXJ")
                  or sheet.cell(row=row, column=1).value == "ZQXJ"]  # fmt: skip
        sheet.cell(row=row, column=4).value = 130
        workbook.save(spreadsheet)
        saved = datetime(2026, 9, 14, 9).timestamp()  # typed in that morning, before opening the app
        os.utime(spreadsheet, (saved, saved))

    def yahoo_finds_zqxj() -> None:
        provider.add("ZQXJ.DE", "EUR", price="91", ranges={date(2026, 9, 4): (Decimal(49), Decimal(51))})

    open_the_app(date(2026, 9, 1))  # the spreadsheet is created
    open_the_app(date(2026, 9, 2))
    open_the_app(date(2026, 9, 3))
    open_the_app(date(2026, 9, 3))  # again the same day: the day's row is updated, not doubled
    open_the_app(date(2026, 9, 7))  # ZQXJ bought: Yahoo does not know it, so the row leaves it out
    open_the_app(date(2026, 9, 8), after=enter_value)  # a value entered today: today's row includes it
    open_the_app(date(2026, 9, 9))  # that value is from an earlier day now: no row
    open_the_app(date(2026, 9, 10), offline=True)  # Yahoo unreachable: last known prices, no row
    open_the_app(date(2026, 9, 11), offline=True)
    open_the_app(date(2026, 9, 14), before=type_into_the_sheet)  # typed in today: today's row includes it
    open_the_app(date(2026, 9, 21), before=yahoo_finds_zqxj)  # Yahoo's price replaces the value typed in
    open_the_app(date(2026, 9, 22), offline=True)  # ZQXJ's last price is Yahoo's now, not the value typed in
    open_the_app(date(2026, 9, 23))

    sheet = load_workbook(spreadsheet)["Portfolio History"]
    written = [cell.value.date() for cell in sheet["A"][1:]]
    assert written == [date(2026, 9, day) for day in (1, 2, 3, 7, 8, 14, 21, 23)]  # one row per valid day
    partial = {when for when, cell in zip(written, sheet["A"][1:]) if cell.comment and "Leaves out ZQXJ" in cell.comment.text}
    assert partial == {date(2026, 9, 7)}
    # Market value on the first day: 10 ABCD at 120 and 4 WXYZ at 60; on the last, 12 ABCD, 2 WXYZ and 2 ZQXJ.
    assert history_as_shown[date(2026, 9, 1)][2] == 1440
    assert history_as_shown[date(2026, 9, 23)][2] == 12 * 142 + 2 * 82 + 2 * 92
    assert history_as_shown[date(2026, 9, 8)][2] == 10 * 127 + 4 * 67 + 120  # with the value entered
    assert history_as_shown[date(2026, 9, 14)][2] == 12 * 133 + 4 * 73 + 130  # with the value typed in

    assert worth(date(2026, 9, 9), "ZQXJ").endswith(YOUR_VALUE)
    assert worth(date(2026, 9, 10), "ABCD").endswith(LAST_KNOWN) and worth(date(2026, 9, 10), "ZQXJ").endswith(YOUR_VALUE)
    assert amount(worth(date(2026, 9, 14), "ZQXJ"), "130.00") and worth(date(2026, 9, 14), "ZQXJ").endswith(YOUR_VALUE)
    assert amount(worth(date(2026, 9, 21), "ZQXJ"), "182.00") and not worth(date(2026, 9, 21), "ZQXJ").endswith(YOUR_VALUE)
    assert amount(worth(date(2026, 9, 22), "ZQXJ"), "182.00") and worth(date(2026, 9, 22), "ZQXJ").endswith(LAST_KNOWN)

    backups = sorted(path.name for path in tmp_path.glob("mine_backup_*.xlsx"))
    assert backups == ["mine_backup_2026-09-21.xlsx", "mine_backup_2026-09-22.xlsx", "mine_backup_2026-09-23.xlsx"]


# -- upgrading ----------------------------------------------------------------------------------------------

# What earlier versions of the app saved, exactly as each wrote it. Newer versions only ever added keys, and
# one ("leave_out_unpriced") was dropped again.
SETTINGS_OF = {
    "2026-10-11": {"ticker_choices": {}, "suffixes": [".DE", ".L", ".AS", ".MI", ".PA"], "index_names": {},
                   "last_folder": None, "spreadsheet_path": None, "holdings_sheet": "Current Holdings",
                   "history_sheet": "Portfolio History"},
    "2026-10-14": {"ticker_choices": {}, "ask_again": [], "always_use_first_listing": False,
                   "leave_out_unpriced": True, "suffixes": [".DE", ".L", ".AS", ".MI", ".PA"], "index_names": {},
                   "last_folder": None, "spreadsheet_path": None, "holdings_sheet": "Current Holdings",
                   "history_sheet": "Portfolio History"},
    "2026-10-15": {"ticker_choices": {}, "ask_again": [], "always_use_first_listing": False,
                   "suffixes": [".DE", ".L", ".AS", ".MI", ".PA"], "index_names": {}, "last_folder": None,
                   "spreadsheet_path": None, "holdings_sheet": "Current Holdings",
                   "history_sheet": "Portfolio History", "backups_kept": 2},
}  # fmt: skip
LAST_PRICE = {"currency": "EUR", "price": "400", "quoted_currency": "EUR", "converted": False,
              "as_of": "2026-10-09T17:35:00+00:00"}  # fmt: skip
ENTERED = {"currency": "EUR", "price": "450", "entered_at": "2026-10-10T12:00:00+00:00"}
PRICES_OF = {
    "2026-10-14 (last prices)": {"prices": {"EQQB.DE": LAST_PRICE}},
    "2026-10-14 (values entered)": {"prices": {"EQQB.DE": LAST_PRICE}, "entered": {"EQQB": ENTERED}},
    "2026-10-15": {"prices": {"EQQB.DE": LAST_PRICE}, "entered": {"EQQB": ENTERED},
                   "written": {"/somewhere/else.xlsx": {"EQQB": "662.98"}}},
}  # fmt: skip


@pytest.mark.parametrize("prices_from", list(PRICES_OF))
@pytest.mark.parametrize("settings_from", list(SETTINGS_OF))
def test_what_an_earlier_version_kept_is_used_after_upgrading(
    app, tmp_path: Path, settings_from: str, prices_from: str
) -> None:
    # Spec J6: installing a newer version over an older one keeps everything.
    import json

    spreadsheet = tmp_path / "mine.xlsx"
    Workbook().save(spreadsheet)  # the user's own spreadsheet, chosen with the earlier version
    saved = dict(SETTINGS_OF[settings_from])
    saved["ticker_choices"] = {
        "EQQB": {"symbol": "EQQB.DE", "name": "Nasdaq fund", "exchange": "XETRA", "currency": "EUR", "confirmed": True}
    }
    saved["index_names"] = {"EQQB": "My Nasdaq"}
    saved["last_folder"], saved["spreadsheet_path"] = str(SAMPLE.parent), str(spreadsheet)
    (tmp_path / "settings.json").write_text(json.dumps(saved, indent=2), encoding="utf-8")
    (tmp_path / "last_prices.json").write_text(json.dumps(PRICES_OF[prices_from], indent=2), encoding="utf-8")

    window, errors = make_window(tmp_path, FakeProvider(offline=True))  # Yahoo unreachable: what is kept shows
    window.ask_listing = lambda resolution: pytest.fail("the listing chosen before is asked for again")

    assert window.spreadsheet_badge.text() == "✓ Already done: remembered from last time."
    assert shows(window.spreadsheet_label, spreadsheet)
    kept = saved.get("backups_kept", 3)  # versions before it was a setting: the default
    assert f"the {kept} most recent are kept" in window.backup_label.text()
    drop(window, SAMPLE)
    run(app, window)

    assert errors == []
    name, worth = cells(window.table)[0][0], cells(window.table)[0][3]
    assert name == "(EQQB) My Nasdaq"
    if "entered" in PRICES_OF[prices_from]:
        assert amount(worth, "662.98") and worth.endswith(YOUR_VALUE)  # 1.47328346 shares at 450
    else:
        assert amount(worth, "589.31") and worth.endswith(LAST_KNOWN)  # at the last price found, 400
    assert load_workbook(spreadsheet)["Current Holdings"]["A2"].value == "(EQQB) My Nasdaq"
    # And what it saves now is read back the same way.
    again = Settings.load(tmp_path / "settings.json")
    assert again.ticker_choices["EQQB"] == TickerChoice("EQQB.DE", "Nasdaq fund", "XETRA", "EUR", confirmed=True)
    assert again.index_names == {"EQQB": "My Nasdaq"} and again.backups_kept == kept


# -- J4: the packaged app's own check ----------------------------------------------------------------------


def test_packaging_smoke_test_reads_the_sample_and_writes_a_spreadsheet(app, tmp_path: Path) -> None:
    # Spec J4: as CI starts the packaged app.
    import json

    from statement_tracker.__main__ import main

    out = tmp_path / "report.json"

    assert main(["--smoke-test", str(SAMPLE), "--smoke-out", str(out)]) == 0
    report = json.loads(out.read_text(encoding="utf-8"))
    assert report["ok"] and report["spreadsheet_written"]
    assert report["sample_rows"] == 18
    assert report["imports"]["openpyxl"]["version"]
