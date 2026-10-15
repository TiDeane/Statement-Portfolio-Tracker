"""Excel export: our two sheets are written, everything else in the workbook is left alone."""

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from fakes import FakeProvider, computed, statement_from, write_statement
from openpyxl import Workbook, load_workbook

from statement_tracker import cli
from statement_tracker.core.calculator import compute_portfolio, value_portfolio
from statement_tracker.core.errors import UserFacingError
from statement_tracker.core.exporters.excel import (
    HISTORY_HEADER,
    HOLDINGS_HEADER,
    MISSING,
    ExcelExporter,
    typed_values,
)
from statement_tracker.core.models import MarketPrice, PortfolioValuation
from statement_tracker.settings import Settings

LINES = (
    "2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,EUR 2000,EUR,1",
    "2025-01-06T10:00:00.000Z,ABCD,BUY - MARKET,10,EUR 100.00,EUR 1001,EUR,1",
    "2025-01-20T10:00:00.000Z,WXYZ,BUY - MARKET,4,EUR 50.00,EUR 201,EUR,1",
)
DAY_ONE, DAY_TWO = date(2026, 10, 12), date(2026, 10, 13)


def valuation(abcd: str | None = "120", wxyz: str | None = "100") -> PortfolioValuation:
    prices = {
        ticker: MarketPrice(Decimal(price), f"{ticker}.DE", "EUR")
        for ticker, price in (("ABCD", abcd), ("WXYZ", wxyz))
        if price is not None
    }
    return value_portfolio(compute_portfolio(statement_from(*LINES)), prices)


def rows(path: Path, sheet: str) -> list[tuple]:
    return [tuple(row) for row in load_workbook(path)[sheet].iter_rows(values_only=True)]


def approx(number: float) -> object:
    return pytest.approx(number, abs=1e-6)


def test_missing_workbook_is_created_with_both_sheets(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"

    result = ExcelExporter(path).export(valuation(), DAY_ONE)

    assert (result.created, result.holdings_written, result.history) == (True, 2, "added")
    assert load_workbook(path).sheetnames == ["Current Holdings", "Portfolio History"]
    holdings = computed(path, "Current Holdings")
    assert tuple(holdings[0]) == HOLDINGS_HEADER
    # Everything worked out from the values is a formula; percentages are stored the Excel way (83.28% is 0.8328).
    # Columns: Investment, Invested, % of Invested, Current Value, % of Current Value, % Up/Down, Dividends, Currency.
    assert holdings[1] == [
        "ABCD", 1001, approx(0.832779), 1200, approx(0.75), approx(0.198801), 0, "EUR",
    ]  # fmt: skip
    assert holdings[2][:5] == ["WXYZ", 201, approx(0.167221), 400, approx(0.25)]
    assert holdings[3] == ["Total", 1202, 1, 1600, 1, approx(0.331115), 0, "EUR"]
    assert len(holdings) == 4
    assert rows(path, "Current Holdings")[1][3] == 1200  # the values themselves are numbers, which can be typed over
    history = rows(path, "Portfolio History")
    assert history[0] == HISTORY_HEADER
    assert history[1][0].date() == DAY_ONE
    # cost basis, market value, unrealized, realized, ROI %, net worth, net deposits, gain vs deposits
    assert history[1][1:] == (1202, 1600, 398, 0, pytest.approx(0.331115, abs=1e-6), 2398, 2000, 398)


def test_sheets_are_formatted_for_reading(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    ExcelExporter(path).export(valuation(), DAY_ONE)
    workbook = load_workbook(path)

    for name, header in (("Current Holdings", HOLDINGS_HEADER), ("Portfolio History", HISTORY_HEADER)):
        sheet = workbook[name]
        first_header, first_value = sheet["A1"], sheet["B2"]
        # Header: its own colour, bold. Data cells: no fill.
        assert first_header.fill.fill_type == "solid" and first_header.fill.start_color.rgb.endswith("1F4E78")
        assert first_header.font.bold
        assert first_value.fill.fill_type is None
        # Every cell in the table has a thin black outline, including empty ones.
        for row in sheet.iter_rows():
            for cell in row:
                for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom):
                    assert side.style == "thin" and side.color.rgb.endswith("000000")
        # Columns are wider than the longest text in them.
        for index, title in enumerate(header):
            width = sheet.column_dimensions[chr(ord("A") + index)].width
            assert width > len(title)

    holdings, history = workbook["Current Holdings"], workbook["Portfolio History"]
    assert [holdings.cell(row=2, column=column).number_format for column in (3, 5, 6)] == ["0.00%"] * 3
    # Rows are taller than Excel's default of 15 points, the header tallest.
    assert holdings.row_dimensions[1].height > holdings.row_dimensions[2].height > 15
    assert all(holdings.row_dimensions[row].height > 15 for row in range(1, 5))
    assert history.row_dimensions[2].height > 15
    assert holdings["B2"].number_format == "0.00"
    assert history["F2"].number_format == "0.00%"
    assert history["A2"].number_format == "yyyy-mm-dd"
    assert holdings["A4"].value == "Total" and holdings["A4"].font.bold and not holdings["A3"].font.bold


def test_share_columns_say_what_they_are_a_share_of_and_older_headers_are_upgraded(tmp_path: Path) -> None:
    assert HOLDINGS_HEADER[1:5] == ("Invested", "% of Invested", "Current Value", "% of Current Value")

    path = tmp_path / "portfolio.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Current Holdings"
    sheet.append(["Index", "Invested", "Current Value", "Current %", "Invested %", "% Up/Down", "Dividends", "Currency"])
    workbook.save(path)

    ExcelExporter(path).export(valuation(), DAY_ONE)

    assert rows(path, "Current Holdings")[0] == HOLDINGS_HEADER


def test_the_first_column_is_called_investment_and_a_sheet_calling_it_index_is_upgraded(tmp_path: Path) -> None:
    assert HOLDINGS_HEADER[0] == "Investment"  # what each row is; an index is what some funds follow
    path = tmp_path / "portfolio.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Current Holdings"
    sheet.append(["Index", *HOLDINGS_HEADER[1:]])  # as this app wrote it until now
    workbook.save(path)

    ExcelExporter(path).export(valuation(), DAY_ONE)

    assert rows(path, "Current Holdings")[0] == HOLDINGS_HEADER


def test_holdings_sheet_from_an_earlier_version_is_rewritten_in_the_shorter_layout(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    earlier = (
        "Index", "Invested", "Current Value", "Current %", "Invested %", "Shares", "Average Cost", "Current Price",
        "Unrealized PnL", "% Up/Down", "Dividends", "Yahoo Symbol", "Price Time", "Currency",
    )  # fmt: skip
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Current Holdings"
    sheet.append(earlier)
    sheet.append(["OLD", 1, 2, 0.5, 0.5, 3, 4, 5, 6, 0.1, 0, "OLD.DE", "2026-10-01T10:00", "EUR"])
    sheet.column_dimensions["N"].width = 40
    workbook.save(path)

    ExcelExporter(path).export(valuation(), DAY_ONE)

    holdings = rows(path, "Current Holdings")
    assert holdings[0] == HOLDINGS_HEADER
    assert [row[0] for row in holdings] == ["Investment", "ABCD", "WXYZ", "Total"]
    assert all(len(row) == len(HOLDINGS_HEADER) for row in holdings)
    assert "N" not in load_workbook(path)["Current Holdings"].column_dimensions


def test_second_run_on_the_same_day_updates_that_days_row(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    ExcelExporter(path).export(valuation(), DAY_ONE)

    result = ExcelExporter(path).export(valuation(abcd="130"), DAY_ONE)

    history = rows(path, "Portfolio History")
    assert result.history == "updated"
    assert len(history) == 2
    assert history[1][2] == 1700
    assert len(rows(path, "Current Holdings")) == 4  # header, two holdings, total


def test_a_new_day_adds_a_row_and_keeps_the_old_ones(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    ExcelExporter(path).export(valuation(), DAY_ONE)

    result = ExcelExporter(path).export(valuation(abcd="130"), DAY_TWO)

    history = rows(path, "Portfolio History")
    assert result.history == "added"
    assert [row[0].date() for row in history[1:]] == [DAY_ONE, DAY_TWO]
    assert (history[1][2], history[2][2]) == (1600, 1700)


def test_holdings_sheet_is_replaced_not_appended_to(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    ExcelExporter(path).export(valuation(), DAY_ONE)
    only_abcd = value_portfolio(
        compute_portfolio(statement_from(*LINES[:2])), {"ABCD": MarketPrice(Decimal("120"), "ABCD.DE", "EUR")}
    )

    ExcelExporter(path).export(only_abcd, DAY_TWO)

    assert [row[0] for row in rows(path, "Current Holdings")] == ["Investment", "ABCD", "Total"]


def _users_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "investimentos"
    sheet.append(["Index", "Investidos€", "ValoresAprox€", "Current%", "Target%"])
    sheet.append(["ABCD", 1000, 1100, "=C2/SUM(C2:C3)", 0.6])
    sheet.append(["WXYZ", 200, 300, "=C3/SUM(C2:C3)", 0.4])
    sheet["J1"] = "Legend"
    workbook.save(path)


def test_existing_workbook_keeps_its_own_sheet_and_gets_a_backup(tmp_path: Path) -> None:
    path = tmp_path / "mine.xlsx"
    _users_workbook(path)
    before = rows(path, "investimentos")

    result = ExcelExporter(path).export(valuation(), DAY_ONE)
    ExcelExporter(path).export(valuation(), DAY_TWO)

    assert result.created is False
    assert load_workbook(path).sheetnames == ["investimentos", "Current Holdings", "Portfolio History"]
    assert rows(path, "investimentos") == before
    assert before[1][3] == "=C2/SUM(C2:C3)"  # formulas survive as formulas
    backup = tmp_path / "mine_backup_before_first_update.xlsx"
    assert "mine_backup_before_first_update.xlsx" in " ".join(result.notes)
    assert load_workbook(backup).sheetnames == ["investimentos"]


def test_a_sheet_with_our_name_that_is_not_ours_is_never_overwritten(tmp_path: Path) -> None:
    path = tmp_path / "mine.xlsx"
    workbook = Workbook()
    workbook.active.title = "Current Holdings"
    workbook.active.append(["My own notes", "do not touch"])
    workbook.save(path)

    with pytest.raises(UserFacingError) as error:
        ExcelExporter(path).export(valuation(), DAY_ONE)

    assert "was not made by this app" in str(error.value)
    assert rows(path, "Current Holdings") == [("My own notes", "do not touch")]
    assert load_workbook(path).sheetnames == ["Current Holdings"]


def test_without_any_price_unavailable_is_written_and_history_is_left_alone(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    ExcelExporter(path).export(valuation(), DAY_ONE)

    result = ExcelExporter(path).export(valuation(abcd=None, wxyz=None), DAY_TWO)

    holdings = computed(path, "Current Holdings")
    assert result.history == "skipped"
    assert holdings[1][:6] == ["ABCD", 1001, approx(0.832779), "unavailable", "unavailable", "unavailable"]
    assert holdings[3][:6] == ["Total", 1202, 1, "unavailable", "unavailable", "unavailable"]
    assert len(rows(path, "Portfolio History")) == 2
    assert "no row was written" in " ".join(result.notes)


def test_last_known_price_is_written_with_a_comment_and_leaves_history_alone(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    ExcelExporter(path).export(valuation(), DAY_ONE)
    portfolio = compute_portfolio(statement_from(*LINES))
    prices = {
        "ABCD": MarketPrice(Decimal("120"), "ABCD.DE", "EUR"),
        "WXYZ": MarketPrice(
            Decimal("100"), "WXYZ.DE", "EUR", as_of=datetime(2026, 10, 9, 17, 35, tzinfo=UTC), last_known=True
        ),
    }

    result = ExcelExporter(path).export(value_portfolio(portfolio, prices), DAY_TWO)

    sheet = load_workbook(path)["Current Holdings"]
    assert [row[3] for row in computed(path, "Current Holdings")[1:]] == [1200, 400, 1600]
    assert sheet["D2"].comment is None
    assert "Last known price, from 2026-10-09" in sheet["D3"].comment.text
    assert "not up to date" in sheet["D4"].comment.text
    assert result.history == "skipped"
    assert len(rows(path, "Portfolio History")) == 2
    assert "Yahoo could not be reached, so no row was written" in " ".join(result.notes)

    ExcelExporter(path).export(valuation(), DAY_TWO)

    sheet = load_workbook(path)["Current Holdings"]
    assert sheet["D3"].comment is None and sheet["D4"].comment is None


def test_left_out_investment_is_greyed_and_the_total_and_history_cover_the_rest(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    portfolio = compute_portfolio(statement_from(*LINES))
    prices = {"ABCD": MarketPrice(Decimal("120"), "ABCD.DE", "EUR")}

    result = ExcelExporter(path).export(value_portfolio(portfolio, prices), DAY_ONE)

    holdings = computed(path, "Current Holdings")
    assert holdings[1][:5] == ["ABCD", 1001, 1, 1200, 1]
    assert holdings[2][:6] == ["WXYZ", 201, "not in total", "unavailable", "not in total", "unavailable"]
    assert holdings[3][:6] == ["Total", 1001, 1, 1200, 1, approx(0.198801)]
    sheet = load_workbook(path)["Current Holdings"]
    # While a row has no value, Excel itself shows it in grey, with its Current Value cell in red
    # to say a value can be typed there: rules, not fixed colours, so typing a value undoes them.
    rules = {str(cells.sqref): cells.rules[0] for cells in sheet.conditional_formatting}
    assert set(rules) == {"A2:H3", "D2:D3"}
    assert rules["A2:H3"].formula == ["NOT(ISNUMBER($D2))"] and rules["A2:H3"].dxf.font.color.rgb.endswith("808080")
    assert rules["D2:D3"].formula == ["NOT(ISNUMBER($D2))"] and rules["D2:D3"].dxf.fill.fgColor.rgb.endswith("FFC7CE")
    assert rules["D2:D3"].priority < rules["A2:H3"].priority  # its dark red text wins over the row's grey
    assert "type this investment's current value here" in sheet["D3"].comment.text
    assert result.history == "added"
    history = rows(path, "Portfolio History")[1]
    assert history[1:4] == (1001, 1200, 199)
    assert (history[6], history[8]) == ("unavailable", "unavailable")  # net worth needs every value
    assert "leaves out what has no price (WXYZ)" in " ".join(result.notes)


def test_a_value_typed_into_the_sheet_updates_its_totals_and_is_read_back(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    portfolio = compute_portfolio(statement_from(*LINES))
    ExcelExporter(path).export(value_portfolio(portfolio, {"ABCD": MarketPrice(Decimal("120"), "ABCD.DE", "EUR")}), DAY_ONE)
    assert typed_values(path) == {}  # "unavailable" is not a value

    # The user types WXYZ's current value over "unavailable", as in Excel.
    workbook = load_workbook(path)
    workbook["Current Holdings"]["D3"] = 400
    workbook.save(path)

    holdings = computed(path, "Current Holdings")
    assert holdings[2][:6] == ["WXYZ", 201, approx(0.167221), 400, approx(0.25), approx(0.990050)]
    assert holdings[3][:6] == ["Total", 1202, 1, 1600, 1, approx(0.331115)]  # everything follows
    assert typed_values(path) == {"WXYZ": (Decimal(400), MISSING)}  # and the app finds it next time


def entered_on(day: date) -> MarketPrice:
    """WXYZ at 100, typed in by the user at noon (local time) on ``day``."""
    noon = datetime(day.year, day.month, day.day, 12).astimezone()
    return MarketPrice(Decimal("100"), "", "EUR", as_of=noon, entered=True)


def test_price_entered_today_is_written_with_a_comment_and_counts_for_history(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    portfolio = compute_portfolio(statement_from(*LINES))
    prices = {"ABCD": MarketPrice(Decimal("120"), "ABCD.DE", "EUR"), "WXYZ": entered_on(DAY_TWO)}

    result = ExcelExporter(path).export(value_portfolio(portfolio, prices), DAY_TWO)

    sheet = load_workbook(path)["Current Holdings"]
    assert [row[3] for row in computed(path, "Current Holdings")[2:]] == [400, 1600]
    assert sheet["D2"].comment is None
    assert "entered by hand on 2026-10-13" in sheet["D3"].comment.text
    assert "Includes prices entered by hand" in sheet["D4"].comment.text
    assert result.history == "added"
    assert rows(path, "Portfolio History")[1][2] == 1600


def test_price_entered_on_an_earlier_day_leaves_history_alone(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    portfolio = compute_portfolio(statement_from(*LINES))
    prices = {"ABCD": MarketPrice(Decimal("120"), "ABCD.DE", "EUR"), "WXYZ": entered_on(DAY_ONE)}

    result = ExcelExporter(path).export(value_portfolio(portfolio, prices), DAY_TWO)

    assert result.history == "skipped"
    assert len(rows(path, "Portfolio History")) == 1
    assert "a price you entered on an earlier day, so no row was written" in " ".join(result.notes)


def test_workbook_open_in_excel_gives_a_friendly_message(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "portfolio.xlsx"
    ExcelExporter(path).export(valuation(), DAY_ONE)
    before = path.read_bytes()

    def locked(source, target) -> None:
        raise PermissionError(13, "The process cannot access the file because it is being used")

    monkeypatch.setattr("os.replace", locked)
    with pytest.raises(UserFacingError) as error:
        ExcelExporter(path).export(valuation(abcd="130"), DAY_TWO)

    assert "open in another program" in str(error.value)
    assert "Close the spreadsheet in Excel" in str(error.value)
    assert path.read_bytes() == before
    assert sorted(item.name for item in tmp_path.iterdir()) == ["portfolio.xlsx"]


@pytest.mark.parametrize("name", ["portfolio.xls", "portfolio.csv", "portfolio.xlsm"])
def test_other_file_types_are_refused(tmp_path: Path, name: str) -> None:
    with pytest.raises(UserFacingError) as error:
        ExcelExporter(tmp_path / name).export(valuation(), DAY_ONE)

    assert ".xlsx" in str(error.value)


def test_file_that_is_not_a_workbook_is_explained_and_untouched(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    path.write_text("not a workbook", encoding="utf-8")

    with pytest.raises(UserFacingError) as error:
        ExcelExporter(path).export(valuation(), DAY_ONE)

    assert "could not be opened as an Excel workbook" in str(error.value)
    assert path.read_text(encoding="utf-8") == "not a workbook"


def test_sheet_names_can_be_changed(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"

    ExcelExporter(path, holdings_sheet="Posições", history_sheet="Histórico").export(valuation(), DAY_ONE)

    assert load_workbook(path).sheetnames == ["Posições", "Histórico"]


def _run(tmp_path: Path, *arguments: str, today: date = DAY_ONE) -> tuple[int, str]:
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120")
    provider.add("WXYZ.DE", "EUR", price="100")
    statement = write_statement(tmp_path, *LINES)
    errors = io.StringIO()
    code = cli.main(
        [str(statement), "--symbol", "ABCD=ABCD.DE", "--symbol", "WXYZ=WXYZ.DE", "--no-prompt",
         "--settings", str(tmp_path / "settings.json"), *arguments],
        provider, io.StringIO(), errors, today,
    )  # fmt: skip
    return code, errors.getvalue()


def test_cli_updates_the_spreadsheet_and_remembers_where_it_is(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"

    code, errors = _run(tmp_path, "--spreadsheet", str(path))
    assert code == 0
    assert "Spreadsheet created" in errors and "2 row(s) written" in errors
    assert "a new row for 2026-10-12 was added" in errors
    assert Settings.load(tmp_path / "settings.json").spreadsheet_path == path.resolve()

    # From then on the remembered workbook is updated without any option.
    code, errors = _run(tmp_path, today=DAY_TWO)
    assert code == 0
    assert "Spreadsheet updated" in errors
    assert len(rows(path, "Portfolio History")) == 3


def test_index_column_shows_the_name_next_to_the_ticker_when_one_is_known(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"

    ExcelExporter(path, index_names={"ABCD": "S&P 500"}).export(valuation(), DAY_ONE)

    assert [row[0] for row in rows(path, "Current Holdings")] == ["Investment", "(ABCD) S&P 500", "WXYZ", "Total"]
    assert load_workbook(path)["Current Holdings"].column_dimensions["A"].width > len("(ABCD) S&P 500")


def test_cli_uses_the_fund_name_from_yahoo_until_the_user_gives_their_own(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"

    _run(tmp_path, "--spreadsheet", str(path))
    assert [row[0] for row in rows(path, "Current Holdings")][1:3] == ["(ABCD) ABCD.DE fund", "(WXYZ) WXYZ.DE fund"]

    _run(tmp_path, "--name", "ABCD=S&P 500")
    assert [row[0] for row in rows(path, "Current Holdings")][1:3] == ["(ABCD) S&P 500", "(WXYZ) WXYZ.DE fund"]

    # The name is remembered for later runs.
    _run(tmp_path, today=DAY_TWO)
    assert rows(path, "Current Holdings")[1][0] == "(ABCD) S&P 500"
    assert Settings.load(tmp_path / "settings.json").index_names == {"ABCD": "S&P 500"}


def test_cli_can_skip_the_remembered_spreadsheet_for_one_run(tmp_path: Path) -> None:
    path = tmp_path / "portfolio.xlsx"
    _run(tmp_path, "--spreadsheet", str(path))
    before = path.read_bytes()

    code, errors = _run(tmp_path, "--no-spreadsheet", today=DAY_TWO)

    assert code == 0
    assert "Spreadsheet" not in errors
    assert path.read_bytes() == before


def test_cli_without_a_chosen_spreadsheet_only_prints_and_says_how_to_choose_one(tmp_path: Path) -> None:
    code, errors = _run(tmp_path)

    assert code == 0
    assert "no spreadsheet was updated" in errors
    assert "--spreadsheet FILE.xlsx" in errors
    assert [item.name for item in tmp_path.glob("*.xlsx")] == []
