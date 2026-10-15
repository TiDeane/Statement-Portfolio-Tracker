"""The UI-free pipeline and the command line, with a mocked provider and a temporary settings file."""

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from fakes import SAMPLE, FakeProvider, ranges_around, write_statement
from openpyxl import load_workbook

from statement_tracker import cli
from statement_tracker.core.calculator import buys_by_ticker
from statement_tracker.core.errors import UserFacingError
from statement_tracker.core.market_data.pricing import price_in_currency, to_major_units
from statement_tracker.core.market_data.provider import PriceUnavailableError, Quote, SymbolInfo
from statement_tracker.core.parser import load_statement
from statement_tracker.core.ticker_resolver import resolve_ticker
from statement_tracker.settings import Settings, TickerChoice
from statement_tracker.workflow import build_report, values_typed_in, with_entered_prices

EQQB_BUYS = buys_by_ticker(load_statement(SAMPLE))["EQQB"]

TWO_HOLDINGS = (
    "2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,EUR 2000,EUR,1",
    "2025-01-06T10:00:00.000Z,ABCD,BUY - MARKET,10,EUR 100.00,EUR 1001,EUR,1",
    "2025-01-20T10:00:00.000Z,WXYZ,BUY - MARKET,4,EUR 50.00,EUR 201,EUR,1",
)


def sample_provider() -> FakeProvider:
    provider = FakeProvider()
    provider.add("EQQB.DE", "EUR", price="400", ranges=ranges_around(EQQB_BUYS), name="Nasdaq fund", exchange="XETRA")
    return provider


def run_cli(*arguments: str, provider: FakeProvider, settings: Path) -> tuple[int, str, str]:
    output, errors = io.StringIO(), io.StringIO()
    code = cli.main([*arguments, "--settings", str(settings), "--no-prompt"], provider, output, errors)
    return code, output.getvalue(), errors.getvalue()


def test_cli_prints_one_row_per_open_holding_for_the_sample(tmp_path: Path) -> None:
    code, output, errors = run_cli(str(SAMPLE), provider=sample_provider(), settings=tmp_path / "settings.json")

    assert code == 0
    assert output == (
        "Investment,Invested (EUR),% of Invested,Current Value (EUR approx.),% of Current Value,Gain/Loss (EUR),"
        "% Up/Down\n"
        "EQQB,541.49,100.00,589.31,100.00,47.82,8.83\n"  # 1.47328346 shares x 400
        "Total,541.49,100.00,589.31,100.00,47.82,8.83\n"
    )
    assert "EQQB was matched to EQQB.DE (Nasdaq fund, XETRA, EUR)" in errors
    assert "IUSA" not in output  # fully sold


def test_cli_shows_each_holdings_share_of_the_total_current_value(tmp_path: Path) -> None:
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120")
    provider.add("WXYZ.DE", "EUR", price="100")
    statement = write_statement(tmp_path, *TWO_HOLDINGS)

    code, output, _ = run_cli(
        str(statement), "--symbol", "ABCD=ABCD.DE", "--symbol", "WXYZ=WXYZ.DE",
        provider=provider, settings=tmp_path / "settings.json",
    )  # fmt: skip

    assert code == 0
    # 1001 and 201 out of 1202 invested; 1200 and 400 out of a total current value of 1600.
    assert output.splitlines() == [
        "Investment,Invested (EUR),% of Invested,Current Value (EUR approx.),% of Current Value,Gain/Loss (EUR),"
        "% Up/Down",
        "ABCD,1001.00,83.28,1200.00,75.00,199.00,19.88",
        "WXYZ,201.00,16.72,400.00,25.00,199.00,99.00",
        "Total,1202.00,100.00,1600.00,100.00,398.00,33.11",
    ]


def test_cli_details_add_the_columns_needed_to_check_the_numbers(tmp_path: Path) -> None:
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120", as_of=datetime(2026, 10, 9, 15, 35, tzinfo=UTC))
    provider.add("WXYZ.DE", "EUR", price="100")
    statement = write_statement(
        tmp_path,
        *TWO_HOLDINGS,
        "2025-02-03T10:00:00.000Z,ABCD,SELL - MARKET,4,EUR 120.00,EUR 479,EUR,1",
        "2025-02-10T10:00:00.000Z,ABCD,DIVIDEND,,,EUR 3.21,EUR,1",
    )

    code, output, _ = run_cli(
        str(statement), "--details", "--symbol", "ABCD=ABCD.DE", "--symbol", "WXYZ=WXYZ.DE",
        provider=provider, settings=tmp_path / "settings.json",
    )  # fmt: skip

    assert code == 0
    assert output.splitlines() == [
        "Investment,Invested (EUR),% of Invested,Current Value (EUR approx.),% of Current Value,Gain/Loss (EUR),"
        "% Up/Down,Shares,Average cost (EUR),Price used (EUR),Yahoo symbol,Price time,Realized gain (EUR),"
        "Dividends (EUR)",
        "ABCD,600.60,74.93,720.00,64.29,119.40,19.88,6,100.1000,120.0000,ABCD.DE,2026-10-09T15:35+00:00,78.60,3.21",
        "WXYZ,201.00,25.07,400.00,35.71,199.00,99.00,4,50.2500,100.0000,WXYZ.DE,unavailable,0.00,0.00",
        "Total,801.60,100.00,1120.00,100.00,318.40,39.72,,,,,,,",
    ]


def test_cli_table_output_is_aligned(tmp_path: Path) -> None:
    code, output, _ = run_cli(
        str(SAMPLE), "--table", provider=sample_provider(), settings=tmp_path / "settings.json"
    )

    lines = output.splitlines()
    assert code == 0
    assert lines[0].split() == [
        "Investment", "Invested", "(EUR)", "%", "of", "Invested", "Current", "Value", "(EUR", "approx.)",
        "%", "of", "Current", "Value", "Gain/Loss", "(EUR)", "%", "Up/Down",
    ]  # fmt: skip
    assert lines[1].split() == ["EQQB", "541.49", "100.00", "589.31", "100.00", "47.82", "8.83"]
    assert lines[2].split() == ["Total", "541.49", "100.00", "589.31", "100.00", "47.82", "8.83"]
    assert len(lines[0]) == len(lines[1]) == len(lines[2])


def test_missing_price_is_left_out_of_the_totals(tmp_path: Path) -> None:
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120")
    provider.add("WXYZ.DE", "EUR", price=None)
    statement = write_statement(tmp_path, *TWO_HOLDINGS)

    code, output, errors = run_cli(
        str(statement), "--symbol", "ABCD=ABCD.DE", "--symbol", "WXYZ=WXYZ.DE",
        provider=provider, settings=tmp_path / "settings.json",
    )  # fmt: skip

    assert code == 0
    # WXYZ's value is unavailable, so the totals and shares cover ABCD alone: nothing is estimated.
    assert output.splitlines()[1:] == [
        "ABCD,1001.00,100.00,1200.00,100.00,199.00,19.88",
        "WXYZ,201.00,not in total,unavailable,not in total,unavailable,unavailable",
        "Total,1001.00,100.00,1200.00,100.00,199.00,19.88",
    ]
    assert "WXYZ: current value unavailable. No current price for WXYZ.DE. It is left out of the totals." in errors


def test_ambiguous_ticker_is_left_unavailable_without_a_chooser(tmp_path: Path) -> None:
    provider = sample_provider()
    provider.add("EQQB.MI", "EUR", price="401", ranges=ranges_around(EQQB_BUYS))

    code, output, errors = run_cli(str(SAMPLE), provider=provider, settings=tmp_path / "settings.json")

    assert code == 0
    assert output.splitlines()[1] == "EQQB,541.49,100.00,unavailable,unavailable,unavailable,unavailable"
    assert "More than one listing fits EQQB" in errors
    assert "EQQB.DE, EQQB.MI" in errors
    assert not (tmp_path / "settings.json").exists()


def test_chooser_decides_an_ambiguous_ticker_and_the_choice_is_confirmed(tmp_path: Path) -> None:
    provider = sample_provider()
    provider.add("EQQB.MI", "EUR", price="401", ranges=ranges_around(EQQB_BUYS))
    settings = Settings.load(tmp_path / "settings.json")

    report = build_report(SAMPLE, provider, settings, chooser=lambda resolution: "EQQB.MI")

    assert report.valuation.holdings[0].price.symbol == "EQQB.MI"
    saved = Settings.load(tmp_path / "settings.json").ticker_choices["EQQB"]
    assert (saved.symbol, saved.confirmed) == ("EQQB.MI", True)


def test_matched_ticker_is_remembered_so_the_next_run_skips_matching(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    run_cli(str(SAMPLE), provider=sample_provider(), settings=settings_path)

    saved = Settings.load(settings_path).ticker_choices["EQQB"]
    assert (saved.symbol, saved.name, saved.exchange, saved.currency, saved.confirmed) == (
        "EQQB.DE", "Nasdaq fund", "XETRA", "EUR", False,
    )  # fmt: skip

    second = sample_provider()
    code, output, _ = run_cli(str(SAMPLE), provider=second, settings=settings_path)

    assert code == 0
    assert output.splitlines()[1] == "EQQB,541.49,100.00,589.31,100.00,47.82,8.83"
    assert second.calls == [("quote", "EQQB.DE")]


def test_confirmed_choice_is_never_replaced_by_automatic_matching(tmp_path: Path) -> None:
    settings = Settings.load(tmp_path / "settings.json")
    settings.remember("EQQB", TickerChoice("EQQB.MI", confirmed=True))
    provider = sample_provider()
    provider.add("EQQB.MI", "EUR", price="401")

    report = build_report(SAMPLE, provider, settings)

    assert report.valuation.holdings[0].price.symbol == "EQQB.MI"
    assert settings.remember("EQQB", TickerChoice("EQQB.DE", confirmed=False)) is False
    assert settings.ticker_choices["EQQB"].symbol == "EQQB.MI"


def test_always_using_the_first_listing_skips_the_question(tmp_path: Path) -> None:
    settings = Settings.load(tmp_path / "settings.json")
    settings.always_use_first_listing = True
    provider = sample_provider()
    provider.add("EQQB.MI", "EUR", price="401", ranges=ranges_around(EQQB_BUYS))
    asked: list[str] = []

    report = build_report(SAMPLE, provider, settings, chooser=lambda resolution: asked.append(resolution.ticker))

    assert asked == []
    assert report.valuation.holdings[0].price.symbol == "EQQB.DE"
    assert any("more than one listing fits, so the first one was used, as you asked" in note for note in report.notes)
    saved = Settings.load(tmp_path / "settings.json").ticker_choices["EQQB"]
    assert (saved.symbol, saved.confirmed) == ("EQQB.DE", False)  # the app's pick, not the user's

    # A listing that exists but does not fit the purchases is taken too, with a note saying it was not checked.
    far_off = {day: (low + 100, high + 100) for day, (low, high) in ranges_around(EQQB_BUYS).items()}
    unmatched = FakeProvider()
    unmatched.add("EQQB.DE", "EUR", price="400", ranges=far_off, name="Nasdaq fund", exchange="XETRA")
    report = build_report(
        SAMPLE, unmatched, _always_first(tmp_path),
        chooser=lambda resolution: asked.append(resolution.ticker),
    )  # fmt: skip
    assert asked == []
    assert report.valuation.holdings[0].price.symbol == "EQQB.DE"
    assert any("no listing could be matched to your purchases" in note and "Check that it is your fund" in note
               for note in report.notes)  # fmt: skip

    # With no listing found at all there is nothing to take: the user is asked.
    nothing = FakeProvider()
    nothing.add("WXYZ.DE", "EUR", price="100")  # no price history, so not a listing for WXYZ's purchases
    statement = write_statement(tmp_path, *TWO_HOLDINGS)
    build_report(
        statement, nothing, settings, symbol_overrides={"ABCD": "WXYZ.DE"},
        chooser=lambda resolution: asked.append(resolution.ticker),
    )  # fmt: skip
    assert asked == ["WXYZ"]


def _always_first(folder: Path) -> Settings:
    """Fresh settings, in their own file, that always take the first listing."""
    settings = Settings.load(folder / "always-first.json")
    settings.always_use_first_listing = True
    return settings


def test_choosing_again_asks_even_when_matching_is_sure_and_only_once(tmp_path: Path) -> None:
    settings = Settings.load(tmp_path / "settings.json")
    provider = sample_provider()
    provider.add("EQQB.MI", "EUR", price="401")  # no history, so EQQB.DE stays the only match
    build_report(SAMPLE, provider, settings)
    settings.ask_again.add("EQQB")
    asked: list[str] = []

    def choose(resolution) -> str:
        asked.append(resolution.ticker)
        return "EQQB.MI"

    report = build_report(SAMPLE, provider, settings, chooser=choose)
    build_report(SAMPLE, provider, settings, chooser=choose)

    assert asked == ["EQQB"]
    assert report.valuation.holdings[0].price.symbol == "EQQB.MI"
    saved = Settings.load(tmp_path / "settings.json")
    assert (saved.ticker_choices["EQQB"].symbol, saved.ticker_choices["EQQB"].confirmed) == ("EQQB.MI", True)
    assert saved.ask_again == set()


def test_skipping_when_asked_again_keeps_the_listing_used_before(tmp_path: Path) -> None:
    settings = Settings.load(tmp_path / "settings.json")
    build_report(SAMPLE, sample_provider(), settings)
    settings.ask_again.add("EQQB")

    report = build_report(SAMPLE, sample_provider(), settings, chooser=lambda resolution: None)

    assert report.valuation.holdings[0].price.symbol == "EQQB.DE"
    assert Settings.load(tmp_path / "settings.json").ask_again == set()


def test_choosing_again_waits_for_a_run_that_can_ask(tmp_path: Path) -> None:
    settings = Settings.load(tmp_path / "settings.json")
    build_report(SAMPLE, sample_provider(), settings)
    settings.ask_again.add("EQQB")

    report = build_report(SAMPLE, sample_provider(), settings)  # no chooser, like the CLI with --no-prompt

    assert report.valuation.holdings[0].price.symbol == "EQQB.DE"
    assert settings.ask_again == {"EQQB"}


def test_forgetting_drops_listings_and_names_and_keeps_the_rest(tmp_path: Path) -> None:
    settings = Settings.load(tmp_path / "settings.json")
    build_report(SAMPLE, sample_provider(), settings)
    settings.index_names = {"EQQB": "Nasdaq", "OLD": "Sold long ago"}
    settings.spreadsheet_path = tmp_path / "mine.xlsx"
    settings.always_use_first_listing = True

    settings.forget("EQQB")
    assert settings.known_tickers() == ["OLD"]
    settings.forget_all()
    assert settings.known_tickers() == []
    assert settings.spreadsheet_path == tmp_path / "mine.xlsx" and settings.always_use_first_listing


def test_symbol_given_by_the_user_replaces_a_saved_choice(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    provider = sample_provider()
    provider.add("EQQB.MI", "EUR", price="401")
    run_cli(str(SAMPLE), provider=provider, settings=settings_path)

    run_cli(str(SAMPLE), "--symbol", "EQQB=EQQB.MI", provider=provider, settings=settings_path)

    saved = Settings.load(settings_path).ticker_choices["EQQB"]
    assert (saved.symbol, saved.confirmed) == ("EQQB.MI", True)


def test_symbol_that_yahoo_does_not_know_is_not_saved(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"

    code, output, errors = run_cli(
        str(SAMPLE), "--symbol", "EQQB=NOPE.XX", provider=sample_provider(), settings=settings_path
    )

    assert code == 0
    assert output.splitlines()[1] == "EQQB,541.49,100.00,unavailable,unavailable,unavailable,unavailable"
    assert "Yahoo has no data for NOPE.XX" in errors
    assert not settings_path.exists()


def test_unreachable_provider_leaves_values_unavailable(tmp_path: Path) -> None:
    code, output, errors = run_cli(
        str(SAMPLE), provider=FakeProvider(offline=True), settings=tmp_path / "settings.json"
    )

    assert code == 0
    assert output.splitlines()[1] == "EQQB,541.49,100.00,unavailable,unavailable,unavailable,unavailable"
    assert "offline" in errors


def test_last_known_price_is_used_when_yahoo_cannot_be_reached(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    online = sample_provider()
    online.listings["EQQB.DE"].as_of = datetime(2026, 10, 9, 17, 35, tzinfo=UTC)
    build_report(SAMPLE, online, Settings.load(settings_path))

    report = build_report(SAMPLE, FakeProvider(offline=True), Settings.load(settings_path))

    holding = report.valuation.holdings[0]
    assert holding.price is not None and holding.price.last_known
    assert (holding.price.price, holding.price.symbol) == (Decimal("400"), "EQQB.DE")
    assert holding.market_value == Decimal("1.47328346") * 400
    assert report.valuation.uses_last_known_prices
    assert report.valuation.total_market_value is not None
    assert (
        "EQQB: Prices could not be fetched (offline). The last price fetched for EQQB.DE, from 2026-10-09, "
        "is used instead, so this value and the totals are not up to date."
    ) in report.notes


def test_cli_marks_the_time_of_a_last_known_price(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    online = sample_provider()
    online.listings["EQQB.DE"].as_of = datetime(2026, 10, 9, 17, 35, tzinfo=UTC)
    run_cli(str(SAMPLE), provider=online, settings=settings_path)

    code, output, errors = run_cli(
        str(SAMPLE), "--details", provider=FakeProvider(offline=True), settings=settings_path
    )

    assert code == 0
    assert output.splitlines()[1].startswith("EQQB,541.49,100.00,589.31,100.00,47.82,8.83,")
    assert ",2026-10-09T17:35+00:00 (last known)," in output.splitlines()[1]
    assert "is used instead" in errors


def test_a_fresh_price_replaces_the_last_known_one(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    build_report(SAMPLE, sample_provider(), Settings.load(settings_path))
    newer = sample_provider()
    newer.listings["EQQB.DE"].price = Decimal("410")
    build_report(SAMPLE, newer, Settings.load(settings_path))

    report = build_report(SAMPLE, FakeProvider(offline=True), Settings.load(settings_path))

    assert report.valuation.holdings[0].price.price == Decimal("410")


def test_last_known_price_in_another_statement_currency_is_not_used(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    build_report(SAMPLE, sample_provider(), Settings.load(settings_path))
    usd = write_statement(
        tmp_path,
        "2025-01-06T09:00:00.000Z,,CASH TOP-UP,,,USD 2000,USD,1",
        "2025-01-06T10:00:00.000Z,EQQB,BUY - MARKET,1,USD 400.00,USD 401,USD,1",
    )

    report = build_report(usd, FakeProvider(offline=True), Settings.load(settings_path))

    assert report.valuation.holdings[0].market_value is None


def abcd_only() -> FakeProvider:
    """ABCD has a price; WXYZ cannot be found on Yahoo at all."""
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120")
    return provider


def test_entered_price_values_an_investment_yahoo_does_not_have_without_fetching(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    statement = write_statement(tmp_path, *TWO_HOLDINGS)
    provider = abcd_only()
    report = build_report(statement, provider, Settings.load(settings_path), symbol_overrides={"ABCD": "ABCD.DE"})
    assert report.valuation.total_market_value == Decimal(1200)  # ABCD alone: WXYZ is left out
    calls = len(provider.calls)

    entered = with_entered_prices(report, {"WXYZ": Decimal("100")}, Settings.load(settings_path))

    assert len(provider.calls) == calls  # nothing fetched again
    wxyz = entered.valuation.holdings[1]
    assert wxyz.price is not None and wxyz.price.entered and wxyz.market_value == Decimal(400)
    assert entered.valuation.total_market_value == Decimal(1600)
    assert entered.valuation.uses_entered_prices
    assert not any("WXYZ: current value unavailable" in note for note in entered.notes)
    assert any(note.startswith("WXYZ: no Yahoo listing is used for it, so the value you entered on") for note in entered.notes)

    # Kept for the next run: Yahoo still has nothing for WXYZ.
    again = build_report(statement, abcd_only(), Settings.load(settings_path))
    assert again.valuation.holdings[1].price.entered
    assert again.valuation.total_market_value == Decimal(1600)


def test_entered_price_is_used_ahead_of_the_last_known_one_and_dropped_once_yahoo_has_a_price(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    build_report(SAMPLE, sample_provider(), Settings.load(settings_path))  # last known: 400
    offline = build_report(SAMPLE, FakeProvider(offline=True), Settings.load(settings_path))

    entered = with_entered_prices(offline, {"EQQB": Decimal("450")}, Settings.load(settings_path))

    price = entered.valuation.holdings[0].price
    assert price.entered and not price.last_known and price.price == Decimal("450")
    assert price.symbol == "EQQB.DE"
    assert any("Prices could not be fetched (offline). The value you entered on" in note for note in entered.notes)
    assert build_report(SAMPLE, FakeProvider(offline=True), Settings.load(settings_path)).valuation.holdings[0].price.entered

    online = build_report(SAMPLE, sample_provider(), Settings.load(settings_path))
    assert not online.valuation.holdings[0].price.entered
    # Dropped: the next offline run goes back to the last price fetched.
    later = build_report(SAMPLE, FakeProvider(offline=True), Settings.load(settings_path))
    assert later.valuation.holdings[0].price.last_known


def test_a_value_typed_into_the_spreadsheet_is_kept_like_an_entered_price(tmp_path: Path) -> None:
    settings_path, sheet = tmp_path / "settings.json", tmp_path / "mine.xlsx"
    statement = write_statement(tmp_path, *TWO_HOLDINGS)
    args = (str(statement), "--symbol", "ABCD=ABCD.DE", "--spreadsheet", str(sheet))
    run_cli(*args, provider=abcd_only(), settings=settings_path)  # WXYZ: "unavailable" in the sheet
    workbook = load_workbook(sheet)
    workbook["Current Holdings"]["D3"] = 404  # typed by the user: the value of 4 shares
    workbook.save(sheet)

    code, output, _ = run_cli(*args, provider=abcd_only(), settings=settings_path)

    assert code == 0
    assert output.splitlines()[2] == "WXYZ,201.00,16.72,404.00,25.19,203.00,101.00"  # 404 of 1604
    assert load_workbook(sheet)["Current Holdings"]["D3"].value == 404
    assert "entered by hand" in load_workbook(sheet)["Current Holdings"]["D3"].comment.text

    # Kept for the next runs, whatever the sheet says, until Yahoo has a price for WXYZ.
    again = build_report(statement, abcd_only(), Settings.load(settings_path))
    assert again.valuation.holdings[1].price.entered and again.valuation.holdings[1].market_value == 404
    with_yahoo = abcd_only()
    with_yahoo.add("WXYZ.DE", "EUR", price="102")
    code, output, _ = run_cli(*args, "--symbol", "WXYZ=WXYZ.DE", provider=with_yahoo, settings=settings_path)
    assert output.splitlines()[2].startswith("WXYZ,201.00,16.72,408.00,")


@pytest.mark.parametrize("entered_in", ["the app", "the spreadsheet"])
def test_a_value_entered_for_a_listing_not_found_gives_way_once_the_listing_is_found(
    tmp_path: Path, entered_in: str
) -> None:
    from statement_tracker.price_cache import PriceCache
    from statement_tracker.workflow import update_spreadsheet

    settings_path, sheet = tmp_path / "settings.json", tmp_path / "mine.xlsx"
    statement = write_statement(tmp_path, *TWO_HOLDINGS)
    args = (str(statement), "--symbol", "ABCD=ABCD.DE", "--spreadsheet", str(sheet))

    # 1. A run in which WXYZ's listing is not found: unavailable.
    code, output, _ = run_cli(*args, provider=abcd_only(), settings=settings_path)
    assert output.splitlines()[2].startswith("WXYZ,201.00,not in total,unavailable,")

    # 2. A value is written for it: 404 for its 4 shares.
    if entered_in == "the app":
        report = build_report(statement, abcd_only(), Settings.load(settings_path))
        report = with_entered_prices(report, {"WXYZ": Decimal("101")}, Settings.load(settings_path))
        update_spreadsheet(report, Settings.load(settings_path), sheet, date.today())  # as the window does
    else:
        workbook = load_workbook(sheet)
        workbook["Current Holdings"]["D3"] = 404
        workbook.save(sheet)
    code, output, _ = run_cli(*args, provider=abcd_only(), settings=settings_path)
    assert output.splitlines()[2].startswith("WXYZ,201.00,16.72,404.00,")  # used while nothing better exists

    # 3. A later run in which the app finds WXYZ's listing by itself: its purchase fits the prices.
    found = abcd_only()
    found.add("WXYZ.DE", "EUR", price="102", ranges={date(2025, 1, 20): (Decimal("49"), Decimal("51"))})
    code, output, errors = run_cli(*args, provider=found, settings=settings_path)

    # 4. Yahoo's price replaces the value written: in the result, the spreadsheet and what is kept.
    assert code == 0
    assert "WXYZ was matched to WXYZ.DE" in errors
    assert output.splitlines()[2].startswith("WXYZ,201.00,16.72,408.00,")
    cell = load_workbook(sheet)["Current Holdings"]["D3"]
    assert cell.value == 408 and cell.comment is None
    assert PriceCache.beside(settings_path).entered_tickers() == set()

    # Gone for good: should Yahoo fail later, the last price it gave is used, not the old value.
    report = build_report(statement, FakeProvider(offline=True), Settings.load(settings_path))
    wxyz = report.valuation.holdings[1]
    assert wxyz.price.last_known and not wxyz.price.entered and wxyz.market_value == 408


def test_values_the_app_wrote_itself_are_not_taken_for_typed_ones(tmp_path: Path) -> None:
    settings_path, sheet = tmp_path / "settings.json", tmp_path / "mine.xlsx"
    args = (str(SAMPLE), "--spreadsheet", str(sheet))
    run_cli(*args, provider=sample_provider(), settings=settings_path)  # last known: 400 a share
    run_cli(*args, provider=FakeProvider(offline=True), settings=settings_path)  # written with its comment
    assert "Last known price" in load_workbook(sheet)["Current Holdings"]["D2"].comment.text

    # The value the app wrote comes back from the sheet as a float, and is not taken for a typed one.
    assert values_typed_in(Settings.load(settings_path), sheet)[0] == {}
    report = build_report(SAMPLE, FakeProvider(offline=True), Settings.load(settings_path))
    assert report.valuation.holdings[0].price.last_known

    # Typed over it, a different value is.
    workbook = load_workbook(sheet)
    workbook["Current Holdings"]["D2"] = 700
    workbook.save(sheet)
    assert values_typed_in(Settings.load(settings_path), sheet)[0] == {"EQQB": Decimal(700)}


def test_a_spreadsheet_left_behind_does_not_undo_a_newer_value(tmp_path: Path) -> None:
    settings_path, sheet = tmp_path / "settings.json", tmp_path / "mine.xlsx"
    statement = write_statement(tmp_path, *TWO_HOLDINGS)
    args = (str(statement), "--symbol", "ABCD=ABCD.DE", "--spreadsheet", str(sheet))
    run_cli(*args, "--price", "WXYZ=100", provider=abcd_only(), settings=settings_path)
    assert load_workbook(sheet)["Current Holdings"]["D3"].value == 400

    # A newer value, entered while the spreadsheet could not be updated (open in Excel, say).
    run_cli(*args, "--no-spreadsheet", "--price", "WXYZ=50", provider=abcd_only(), settings=settings_path)
    assert load_workbook(sheet)["Current Holdings"]["D3"].value == 400  # the sheet still has the old one

    code, output, _ = run_cli(*args, provider=abcd_only(), settings=settings_path)

    assert code == 0
    assert output.splitlines()[2].startswith("WXYZ,201.00,16.72,200.00,")  # the newer value, not the sheet's
    assert load_workbook(sheet)["Current Holdings"]["D3"].value == 200


def test_a_price_forgotten_in_settings_does_not_come_back_from_the_spreadsheet(tmp_path: Path) -> None:
    from statement_tracker.price_cache import PriceCache

    settings_path, sheet = tmp_path / "settings.json", tmp_path / "mine.xlsx"
    statement = write_statement(tmp_path, *TWO_HOLDINGS)
    args = (str(statement), "--symbol", "ABCD=ABCD.DE", "--spreadsheet", str(sheet))
    run_cli(*args, "--price", "WXYZ=100", provider=abcd_only(), settings=settings_path)
    assert load_workbook(sheet)["Current Holdings"]["D3"].value == 400  # written with "entered by hand"

    prices = PriceCache.beside(settings_path)  # what "Forget all" in Settings does
    prices.forget_all_entered()
    prices.save()
    code, output, _ = run_cli(*args, provider=abcd_only(), settings=settings_path)

    assert code == 0
    assert output.splitlines()[2].startswith("WXYZ,201.00,not in total,unavailable,")
    assert PriceCache.beside(settings_path).entered_tickers() == set()


def test_cli_price_is_used_when_yahoo_has_none(tmp_path: Path) -> None:
    statement = write_statement(tmp_path, *TWO_HOLDINGS)

    code, output, errors = run_cli(
        str(statement), "--details", "--symbol", "ABCD=ABCD.DE", "--price", "WXYZ=100",
        provider=abcd_only(), settings=tmp_path / "settings.json",
    )  # fmt: skip

    assert code == 0
    wxyz = output.splitlines()[2]
    assert wxyz.startswith("WXYZ,201.00,16.72,400.00,25.00,199.00,99.00,4,50.2500,100.0000,none,")
    assert "(entered by hand)" in wxyz
    assert output.splitlines()[3].startswith("Total,1202.00,100.00,1600.00,")
    assert "the value you entered on" in errors


def test_cli_price_is_not_used_when_yahoo_has_one(tmp_path: Path) -> None:
    code, output, errors = run_cli(
        str(SAMPLE), "--price", "EQQB=1", provider=sample_provider(), settings=tmp_path / "settings.json"
    )

    assert code == 0
    assert output.splitlines()[1] == "EQQB,541.49,100.00,589.31,100.00,47.82,8.83"
    assert "Yahoo has a current price for EQQB.DE, so it is used instead of the value you entered" in errors


@pytest.mark.parametrize("price", ["abc", "0", "-5", "1,5", "NaN"])
def test_cli_rejects_a_price_that_is_not_a_positive_number(tmp_path: Path, price: str) -> None:
    code, output, errors = run_cli(
        str(SAMPLE), "--price", f"EQQB={price}", provider=sample_provider(), settings=tmp_path / "settings.json"
    )

    assert (code, output) == (2, "")
    assert "is not a valid price" in errors


def test_a_damaged_last_prices_file_is_ignored(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    (tmp_path / "last_prices.json").write_text("{not json", encoding="utf-8")

    report = build_report(SAMPLE, sample_provider(), Settings.load(settings_path))

    assert report.valuation.holdings[0].market_value is not None
    assert "EQQB.DE" in (tmp_path / "last_prices.json").read_text(encoding="utf-8")


def test_price_in_another_currency_is_converted_and_labelled(tmp_path: Path) -> None:
    provider = FakeProvider()
    provider.add("EQQB", "USD", price="500")
    provider.add("USDEUR=X", "EUR", price="0.9")

    code, output, errors = run_cli(
        str(SAMPLE), "--symbol", "EQQB=EQQB", provider=provider, settings=tmp_path / "settings.json"
    )

    assert code == 0
    assert output.splitlines()[1] == "EQQB,541.49,100.00,662.98,100.00,121.49,22.44"  # 1.47328346 x 500 x 0.9
    assert "quoted in USD and was converted to EUR" in errors


def test_progress_is_reported_for_each_investment(tmp_path: Path) -> None:
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120")
    provider.add("WXYZ.DE", "EUR", price="100")
    seen: list[tuple[int, int, str]] = []

    build_report(
        write_statement(tmp_path, *TWO_HOLDINGS),
        provider,
        Settings.load(tmp_path / "settings.json"),
        symbol_overrides={"ABCD": "ABCD.DE", "WXYZ": "WXYZ.DE"},
        progress=lambda done, total, text: seen.append((done, total, text)),
    )

    # Getting ready is a step of its own; then the bar moves as each price arrives (they are fetched together).
    assert seen == [
        (0, 0, "Reading your statement..."),
        (0, 3, "Getting ready to look up prices..."),
        (1, 3, "Getting the current prices (0 of 2)..."),
        (2, 3, "Getting the current prices (1 of 2)..."),
        (3, 3, "Getting the current prices (2 of 2)..."),
        (3, 3, "Adding everything up..."),
    ]
    assert provider.prepared == 1


def test_each_listing_to_find_is_a_step_of_the_progress_too(tmp_path: Path) -> None:
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120", ranges={date(2025, 1, 6): (Decimal("99"), Decimal("101"))})
    provider.add("WXYZ.DE", "EUR", price="100", ranges={date(2025, 1, 20): (Decimal("49"), Decimal("51"))})
    seen: list[tuple[int, int, str]] = []

    build_report(
        write_statement(tmp_path, *TWO_HOLDINGS),
        provider,
        Settings.load(tmp_path / "settings.json"),
        progress=lambda done, total, text: seen.append((done, total, text)),
    )

    # Getting ready, two listings to find, two prices: five steps, each one moving the bar.
    assert seen == [
        (0, 0, "Reading your statement..."),
        (0, 5, "Getting ready to look up prices..."),
        (1, 5, "Finding the stock exchange listing of ABCD..."),
        (2, 5, "Finding the stock exchange listing of WXYZ..."),
        (3, 5, "Getting the current prices (0 of 2)..."),
        (4, 5, "Getting the current prices (1 of 2)..."),
        (5, 5, "Getting the current prices (2 of 2)..."),
        (5, 5, "Adding everything up..."),
    ]

    # Next time both listings are known: only the prices are left to count.
    seen.clear()
    build_report(
        write_statement(tmp_path, *TWO_HOLDINGS),
        provider,
        Settings.load(tmp_path / "settings.json"),
        progress=lambda done, total, text: seen.append((done, total, text)),
    )
    assert [total for _, total, _ in seen[1:]] == [3] * 5


def test_prices_are_fetched_at_the_same_time(tmp_path: Path) -> None:
    import threading

    both_asking = threading.Barrier(2, timeout=5)

    class WaitsForTheOther(FakeProvider):
        """Each price request waits until the other one has started: only works if they run together."""

        def quote(self, symbol: str):
            both_asking.wait()
            return super().quote(symbol)

    provider = WaitsForTheOther()
    provider.add("ABCD.DE", "EUR", price="120")
    provider.add("WXYZ.DE", "EUR", price="100")

    report = build_report(
        write_statement(tmp_path, *TWO_HOLDINGS),
        provider,
        Settings.load(tmp_path / "settings.json"),
        symbol_overrides={"ABCD": "ABCD.DE", "WXYZ": "WXYZ.DE"},
    )

    assert [holding.market_value for holding in report.valuation.holdings] == [Decimal("1200"), Decimal("400")]


def test_notes_keep_the_statement_order_whatever_order_prices_arrive_in(tmp_path: Path) -> None:
    import time

    class FirstIsSlow(FakeProvider):
        def quote(self, symbol: str):
            if symbol == "ABCD":
                time.sleep(0.2)  # ABCD's price arrives after WXYZ's
            return super().quote(symbol)

    provider = FirstIsSlow()
    provider.add("ABCD", "USD", price="130")
    provider.add("WXYZ", "USD", price="110")
    provider.add("USDEUR=X", "EUR", price="0.9")

    report = build_report(
        write_statement(tmp_path, *TWO_HOLDINGS),
        provider,
        Settings.load(tmp_path / "settings.json"),
        symbol_overrides={"ABCD": "ABCD", "WXYZ": "WXYZ"},
    )

    converted = [note for note in report.notes if "was converted" in note]
    assert [note.split(":")[0] for note in converted] == ["ABCD", "WXYZ"]
    # One exchange rate for both: fetched once.
    assert [call for call in provider.calls if call == ("quote", "USDEUR=X")] == [("quote", "USDEUR=X")]
    assert [holding.market_value for holding in report.valuation.holdings] == [
        Decimal("1170.0"),  # 10 x 130 x 0.9
        Decimal("396.0"),  # 4 x 110 x 0.9
    ]


def test_nothing_is_prepared_without_investments_to_price(tmp_path: Path) -> None:
    provider = FakeProvider()
    seen: list[tuple[int, int, str]] = []

    build_report(
        write_statement(tmp_path, TWO_HOLDINGS[0]),
        provider,
        Settings.load(tmp_path / "settings.json"),
        progress=lambda done, total, text: seen.append((done, total, text)),
    )

    assert provider.prepared == 0
    assert seen == [(0, 0, "Reading your statement..."), (0, 0, "Adding everything up...")]


def test_old_latest_price_is_pointed_out(tmp_path: Path) -> None:
    now = datetime(2026, 10, 12, 9, 0, tzinfo=UTC)
    settings = Settings.load(tmp_path / "settings.json")
    fresh, stale = sample_provider(), FakeProvider()
    fresh.listings["EQQB.DE"].as_of = datetime(2026, 10, 9, 17, 35, tzinfo=UTC)  # Friday close, read on Monday
    stale.add("EQQB.DU", "EUR", price="400", as_of=datetime(2026, 9, 28, 12, 0, tzinfo=UTC))

    fresh_report = build_report(SAMPLE, fresh, settings, now=now)
    stale_report = build_report(SAMPLE, stale, settings, symbol_overrides={"EQQB": "EQQB.DU"}, now=now)

    assert not any("out of date" in note for note in fresh_report.notes)
    assert "EQQB: the latest price Yahoo has for EQQB.DU is from 2026-09-28" in " ".join(stale_report.notes)
    assert stale_report.valuation.holdings[0].market_value is not None


def test_pence_are_converted_to_pounds() -> None:
    assert to_major_units(Decimal("4120"), "GBp") == (Decimal("41.2"), "GBP")
    assert to_major_units(Decimal("4120"), "GBX") == (Decimal("41.2"), "GBP")
    assert to_major_units(Decimal("41.2"), "GBP") == (Decimal("41.2"), "GBP")

    provider = FakeProvider()
    provider.add("GBPEUR=X", "EUR", price="1.2")
    quote = Quote(SymbolInfo("ABCD.L"), Decimal("4000"), "GBp")

    same = price_in_currency(quote, "GBP", provider)
    converted = price_in_currency(quote, "EUR", provider)

    assert (same.price, same.converted) == (Decimal("40"), False)
    assert (converted.price, converted.converted, converted.quoted_currency) == (Decimal("48.0"), True, "GBp")


def test_missing_exchange_rate_makes_the_price_unavailable() -> None:
    quote = Quote(SymbolInfo("ABCD"), Decimal("500"), "USD")

    with pytest.raises(PriceUnavailableError):
        price_in_currency(quote, "EUR", FakeProvider())


def test_cli_explains_a_wrong_file_without_a_stack_trace(tmp_path: Path) -> None:
    statement = write_statement(tmp_path, "2025-01-06,ABCD,1", header="Date,Symbol,Amount")

    code, output, errors = run_cli(str(statement), provider=FakeProvider(), settings=tmp_path / "settings.json")

    assert code == 2
    assert output == ""
    assert "does not look like a Revolut stock statement" in errors
    assert "Found columns: Date, Symbol, Amount" in errors
    assert "Traceback" not in errors


def test_cli_stops_on_an_unknown_type_unless_told_to_ignore_it(tmp_path: Path) -> None:
    statement = write_statement(
        tmp_path, *TWO_HOLDINGS, "2025-02-03T10:00:00.000Z,ABCD,STOCK SPLIT,10,,EUR 0,EUR,1"
    )
    provider = FakeProvider()
    provider.add("ABCD.DE", "EUR", price="120")
    provider.add("WXYZ.DE", "EUR", price="100")
    symbols = ("--symbol", "ABCD=ABCD.DE", "--symbol", "WXYZ=WXYZ.DE")

    code, output, errors = run_cli(str(statement), *symbols, provider=provider, settings=tmp_path / "s.json")
    assert code == 2
    assert output == ""
    assert 'Row 5: the transaction type "STOCK SPLIT"' in errors

    code, output, errors = run_cli(
        str(statement), *symbols, "--ignore-unknown-types", provider=provider, settings=tmp_path / "s.json"
    )
    assert code == 0
    assert len(output.splitlines()) == 4  # header, two holdings, Total
    assert "RESULTS MAY BE WRONG" in errors


def test_cli_rejects_a_malformed_symbol_option(tmp_path: Path) -> None:
    code, _, errors = run_cli(
        str(SAMPLE), "--symbol", "EQQB", provider=sample_provider(), settings=tmp_path / "settings.json"
    )

    assert code == 2
    assert "TICKER=YAHOOSYMBOL" in errors


@pytest.mark.parametrize(
    ("typed", "expected"),
    [("2\n", "EQQB.MI"), ("EQQB.SW\n", "EQQB.SW"), ("\n", None), ("9\n", None), ("", None)],
)
def test_terminal_prompt_accepts_a_number_a_symbol_or_nothing(
    monkeypatch: pytest.MonkeyPatch, typed: str, expected: str | None
) -> None:
    provider = sample_provider()
    provider.add("EQQB.MI", "EUR", price="401", ranges=ranges_around(EQQB_BUYS), exchange="Milan")
    resolution = resolve_ticker("EQQB", EQQB_BUYS, "EUR", provider)
    monkeypatch.setattr("sys.stdin", io.StringIO(typed))
    shown = io.StringIO()

    assert cli._prompting_chooser(shown)(resolution) == expected
    assert "More than one listing fits EQQB" in shown.getvalue()
    assert "1. EQQB.DE  Nasdaq fund, XETRA, EUR  [2 of 2 of your purchases fit its prices]" in shown.getvalue()
    assert "2. EQQB.MI" in shown.getvalue()


def test_damaged_settings_file_is_explained(tmp_path: Path) -> None:
    settings_path = tmp_path / "settings.json"
    settings_path.write_text("{not json", encoding="utf-8")

    with pytest.raises(UserFacingError) as error:
        Settings.load(settings_path)

    assert "saved settings could not be read" in str(error.value)


def test_the_days_of_backups_kept_are_saved_and_kept_in_range(tmp_path: Path) -> None:
    path = tmp_path / "settings.json"
    assert Settings.load(path).backups_kept == 3
    for saved, loaded in ((0, 0), (7, 7), (-2, 0), (500, 99)):
        settings = Settings(path)
        settings.backups_kept = saved
        settings.save()
        assert Settings.load(path).backups_kept == loaded
