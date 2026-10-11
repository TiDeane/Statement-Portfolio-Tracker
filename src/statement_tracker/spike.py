"""Phase 0 packaging spike helpers. No Qt imports here.

Temporary: Phase 1 replaces this module with the real parser and workflow.
"""

from __future__ import annotations

import csv
import io
import sys
import time
from dataclasses import dataclass
from pathlib import Path

EXPECTED_HEADER: tuple[str, ...] = (
    "Date",
    "Ticker",
    "Type",
    "Quantity",
    "Price per share",
    "Total Amount",
    "Currency",
    "FX Rate",
)


class HeaderMismatchError(ValueError):
    """The file's first row is not the statement header we know."""

    def __init__(self, expected: tuple[str, ...], found: tuple[str, ...]) -> None:
        self.expected = expected
        self.found = found
        super().__init__(
            "This file does not look like a Revolut stock statement.\n"
            f"Expected columns: {', '.join(expected)}\n"
            f"Found columns: {', '.join(found) if found else '(none)'}"
        )


@dataclass(frozen=True)
class ImportReport:
    """How long a heavy dependency took to import, and which version loaded."""

    name: str
    version: str
    seconds: float


def bundled_sample_path() -> Path:
    """Return the sanitized sample statement shipped with the app."""
    if getattr(sys, "frozen", False):
        base = Path(sys._MEIPASS)  # type: ignore[attr-defined]
    else:
        base = Path(__file__).resolve().parents[2]
    return base / "samples" / "revolut_sample.csv"


def count_statement_rows(path: Path) -> int:
    """Validate the header exactly and return the number of data rows."""
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        header = tuple(next(reader, ()))
        if header != EXPECTED_HEADER:
            raise HeaderMismatchError(EXPECTED_HEADER, header)
        return sum(1 for row in reader if any(cell.strip() for cell in row))


def import_yfinance() -> ImportReport:
    """Import yfinance (slow: it pulls in pandas and numpy) and time it."""
    started = time.perf_counter()
    import yfinance

    return ImportReport("yfinance", yfinance.__version__, time.perf_counter() - started)


def import_openpyxl() -> ImportReport:
    """Import openpyxl and time it."""
    started = time.perf_counter()
    import openpyxl

    return ImportReport("openpyxl", openpyxl.__version__, time.perf_counter() - started)


def openpyxl_roundtrip() -> bool:
    """Write a workbook to memory and read it back, to prove openpyxl works when packaged."""
    from openpyxl import Workbook, load_workbook

    buffer = io.BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = "spike"
    workbook.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer).active["A1"].value == "spike"


def fetch_last_price(symbol: str) -> tuple[str, str]:
    """Return (last price, currency) for a Yahoo symbol. Uses the network."""
    import yfinance

    info = yfinance.Ticker(symbol).fast_info
    return str(info["lastPrice"]), str(info["currency"])
