"""Excel export with openpyxl.

Only two sheets are ever written, and only if they are recognisably ours; every
other sheet in the user's workbook is left exactly as it is. The holdings sheet
is rewritten each run; the history sheet gets one row per date.

Saving rewrites the whole file. openpyxl keeps values, formulas, styles,
conditional formatting and charts through that, but drops pictures (without
Pillow), shapes and text boxes, so before the first write of each day a dated
copy of an existing workbook is made next to it; the newest few are kept.
"""

from __future__ import annotations

import os
import re
import shutil
from collections.abc import Mapping
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from statement_tracker import APP_NAME
from statement_tracker.core.errors import UserFacingError
from statement_tracker.core.exporters.base import ExportResult
from statement_tracker.core.models import PortfolioValuation

DEFAULT_HOLDINGS_SHEET = "Current Holdings"
DEFAULT_HISTORY_SHEET = "Portfolio History"
UNAVAILABLE = "unavailable"
NOT_IN_TOTAL = "not in total"
TOTAL_LABEL = "Total"
MISSING_COMMENT = "No current price from Yahoo"
TYPE_A_VALUE = (
    f"{MISSING_COMMENT}. You can type this investment's current value here, for example the one "
    "the Revolut app shows: the totals follow, and the app keeps it next time."
)
LAST_KNOWN_COMMENT = "Last known price"
PARTIAL_COMMENT = "Leaves out"
"""How the comment on a history row that leaves out investments without a value begins."""
_PARTIAL_FILL_COLOUR = "FCE4D6"
_PARTIAL_TEXT_COLOUR = "833C0B"
"""Light and dark orange from Excel's own theme colours, for history rows that leave something out."""
BACKUPS_KEPT = 3
"""Backup copies kept next to the spreadsheet unless the user chooses otherwise, in all: one is made on
each day it is updated, and the oldest is deleted when there are more."""
MAX_BACKUPS_KEPT = 99
ENTERED_COMMENT = "Value entered by hand"
_EARLIER_ENTERED_COMMENT = "Price per share entered by hand"
"""What the comment said in an earlier version, still recognised when a sheet is read back."""
_LEFT_OUT_COLOUR = "808080"
_TYPE_HERE_COLOUR = "FFC7CE"
"""Light red, Excel's own "Bad" cell colour: the value is missing and can be typed in."""
_TYPE_HERE_TEXT_COLOUR = "9C0006"
"""The dark red text Excel puts on that colour."""

HOLDINGS_HEADER = (
    "Investment",
    "Invested",
    "% of Invested",
    "Current Value",
    "% of Current Value",
    "% Up/Down",
    "Dividends",
    "Currency",
)
# Headers this app wrote in earlier versions. A sheet with one of these is still ours,
# so it is rewritten in the current layout instead of being refused.
_EARLIER_HOLDINGS_HEADERS = (
    ("Index", *HOLDINGS_HEADER[1:]),
    (
        "Index",
        "Invested",
        "Current Value",
        "Share of Current Value",
        "Share of Invested",
        "% Up/Down",
        "Dividends",
        "Currency",
    ),
    ("Index", "Invested", "Current Value", "Current %", "Invested %", "% Up/Down", "Dividends", "Currency"),
    (
        "Index",
        "Invested",
        "Current Value",
        "Current %",
        "Invested %",
        "Shares",
        "Average Cost",
        "Current Price",
        "Unrealized PnL",
        "% Up/Down",
        "Dividends",
        "Yahoo Symbol",
        "Price Time",
        "Currency",
    ),
)
HISTORY_HEADER = (
    "Date",
    "Total Cost Basis",
    "Total Market Value",
    "Unrealized PnL",
    "Realized PnL",
    "ROI %",
    "Net Worth",
    "Net Deposits",
    "Gain vs Deposits",
)
_MONEY_FORMAT = "0.00"
_PERCENT_FORMAT = "0.00%"
_DATE_FORMAT = "yyyy-mm-dd"
# Column number -> number format. Columns not listed keep Excel's general format.
_VALUE_COLUMN = HOLDINGS_HEADER.index("Current Value") + 1
_LAST_COLUMN = chr(ord("A") + len(HOLDINGS_HEADER) - 1)
_HOLDINGS_FORMATS = {
    2: _MONEY_FORMAT,
    3: _PERCENT_FORMAT,
    4: _MONEY_FORMAT,
    5: _PERCENT_FORMAT,
    6: _PERCENT_FORMAT,
    7: _MONEY_FORMAT,
}
_HEADER_ROW_HEIGHT = 30
_ROW_HEIGHT = 24
"""Points. Excel's default is 15; taller rows are easier to read."""
_HISTORY_FORMATS = {1: _DATE_FORMAT, 6: _PERCENT_FORMAT} | {
    column: _MONEY_FORMAT for column in (2, 3, 4, 5, 7, 8, 9)
}
_HEADER_FILL_COLOUR = "1F4E78"
_HEADER_TEXT_COLOUR = "FFFFFF"
_BORDER_COLOUR = "000000"
_HUNDRED = Decimal(100)


def _cell(value: Decimal | None) -> Decimal | str:
    """A number for Excel, or the word "unavailable": never a blank that formulas read as zero."""
    return UNAVAILABLE if value is None else value


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def _signed(amount: Decimal) -> str:
    """An amount to add in a formula: "+504.76" or "-3.20"."""
    return f"{amount:+f}"


_NET_WORTH_CASH = re.compile(r"C\d+([+-]\d+(?:\.\d+)?)")
"""The cash in a history row's Net Worth formula, e.g. "+504.76" in '=IF(ISNUMBER(C5),C5+504.76,...)'."""


def _reference(sheet_name: str) -> str:
    """How a formula names another sheet: 'Current Holdings'!"""
    return "'" + sheet_name.replace("'", "''") + "'!"


def _points_elsewhere(value: Any) -> bool:
    """Whether a cell is a formula that reads another sheet."""
    return isinstance(value, str) and value.startswith("=") and "!" in value


def _ticker(label: str) -> str:
    """The ticker in an Investment cell, which reads "(TICKER) Name", or the ticker alone."""
    return (label[1 : label.index(")")] if label.startswith("(") and ")" in label else label).strip()


def _holdings_totals(sheet: Any) -> tuple[Decimal, Decimal | None, list[str]]:
    """What the holdings sheet's Total row shows now, worked out as its formulas do: the invested
    and current value of the investments that have a value (all invested, and no value, when none
    has one), and the tickers left out for having none."""
    invested_all = invested = value_total = Decimal(0)
    valued = False
    missing: list[str] = []
    for row in range(2, sheet.max_row + 1):
        label, cost, value = (sheet.cell(row=row, column=column).value for column in (1, 2, _VALUE_COLUMN))
        if label is None or label == TOTAL_LABEL or not _is_number(cost):
            continue
        invested_all += Decimal(str(cost))
        if _is_number(value):
            valued = True
            invested += Decimal(str(cost))
            value_total += Decimal(str(value))
        else:
            missing.append(_ticker(str(label)))
    return (invested, value_total, missing) if valued else (invested_all, None, [])


def _partial_comment(tickers: list[str], holdings_sheet: str, live: bool) -> str:
    """The comment on a history row that leaves out investments without a value."""
    names = ", ".join(tickers)
    them = "it" if len(tickers) == 1 else "them"
    when = "when this row was written" if live else "on this day"
    text = f"{PARTIAL_COMMENT} {names}, with no current value {when}. "
    while_coloured = "While this row is coloured, its" if live else "Its"
    text += (
        f'Like the Total row in "{holdings_sheet}", it counts only what has a value. {while_coloured} '
        f"Net Worth, Gain vs Deposits and ROI % count {them} as nothing and are lower than they should be."
    )
    if live:
        text += f' Typing the value into "{holdings_sheet}" before you run the app on another day brings {them} in.'
    return text


def _percent(value: Decimal | None) -> Decimal | str:
    """A percentage as Excel stores it (21.08% is 0.2108), shown with a % sign by the cell format."""
    return UNAVAILABLE if value is None else value / _HUNDRED


def _comment(text: str, height: int = 70) -> Any:
    from openpyxl.comments import Comment

    return Comment(text, APP_NAME, width=220, height=height)


def _shown_length(value: Any, number_format: str | None) -> int:
    """Roughly how many characters Excel shows for a cell, for sizing columns."""
    if value is None:
        return 0
    if isinstance(value, (date, datetime)):
        return 10
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        if number_format == _PERCENT_FORMAT:
            return len(f"{value * 100:.2f}%")
        if number_format == _MONEY_FORMAT:
            return len(f"{value:.2f}")
    return len(str(value))


def _style(
    sheet: Any,
    formats: dict[int, str],
    bold_last_row: bool = False,
    shown: Mapping[tuple[int, int], Any] | None = None,
) -> None:
    """Coloured bold header, black outline on every cell, number formats, columns wide enough for their text.

    ``shown`` gives, by (row, column), what a formula cell shows, so that the column is sized by
    that rather than by the formula's text.
    """
    shown = shown or {}
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import column_index_from_string, get_column_letter

    edge = Side(style="thin", color=_BORDER_COLOUR)
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    header_fill = PatternFill(fill_type="solid", start_color=_HEADER_FILL_COLOUR, end_color=_HEADER_FILL_COLOUR)
    header_font = Font(bold=True, color=_HEADER_TEXT_COLOUR)

    for column in range(1, sheet.max_column + 1):
        number_format = formats.get(column)
        widest = 0
        for row in range(1, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column)
            cell.border = border
            if row == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                widest = max(widest, len(str(cell.value or "")))
            else:
                if number_format:
                    cell.number_format = number_format
                cell.font = Font(bold=bold_last_row and row == sheet.max_row)
                cell.alignment = Alignment(vertical="center")
                widest = max(widest, _shown_length(shown.get((row, column), cell.value), number_format))
        sheet.column_dimensions[get_column_letter(column)].width = widest + 3
    for row in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row].height = _HEADER_ROW_HEIGHT if row == 1 else _ROW_HEIGHT
    # Forget widths of columns an earlier, wider layout used.
    for letter in list(sheet.column_dimensions):
        if column_index_from_string(letter) > sheet.max_column:
            del sheet.column_dimensions[letter]
    sheet.freeze_panes = "A2"


def _mark_missing_values(sheet: Any, last_row: int) -> None:
    """Show, in Excel itself, each investment row without a current value: in grey, as it is left out
    of the totals, and with its Current Value cell filled in red, to say a value can be typed there.
    Typing one makes the row look like the others again. Replaces the formatting rules of this sheet."""
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill

    sheet.conditional_formatting = ConditionalFormattingList()  # the app's sheet: no rules but its own
    if last_row < 2:
        return
    missing = "NOT(ISNUMBER($D2))"
    # Added first, so that it comes first for Excel: its text colour wins over the row's grey.
    fill = PatternFill(fill_type="solid", start_color=_TYPE_HERE_COLOUR, end_color=_TYPE_HERE_COLOUR)
    sheet.conditional_formatting.add(
        f"D2:D{last_row}", FormulaRule(formula=[missing], fill=fill, font=Font(color=_TYPE_HERE_TEXT_COLOUR))
    )
    sheet.conditional_formatting.add(
        f"A2:{_LAST_COLUMN}{last_row}", FormulaRule(formula=[missing], font=Font(color=_LEFT_OUT_COLOUR))
    )


MISSING, ENTERED, LAST_KNOWN = "missing", "entered", "last known"
"""What the app had written in a Current Value cell it commented on, as ``typed_values`` reports it."""
_COMMENT_KINDS = (
    (MISSING_COMMENT, MISSING),
    (ENTERED_COMMENT, ENTERED),
    (_EARLIER_ENTERED_COMMENT, ENTERED),
    (LAST_KNOWN_COMMENT, LAST_KNOWN),
)


def typed_values(path: Path, holdings_sheet: str = DEFAULT_HOLDINGS_SHEET) -> dict[str, tuple[Decimal, str]]:
    """Return, by ticker, the current values found in the app's holdings sheet where the app left a
    comment, with what the app had written there: MISSING ("unavailable", so any number there was
    typed by the user), ENTERED (a price the user entered) or LAST_KNOWN (an earlier price).

    For the last two the number may be the app's own; the caller tells them apart. A file that
    cannot be read gives nothing: this must never stop a run.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path)
    except Exception:
        return {}
    if holdings_sheet not in workbook.sheetnames:
        return {}
    sheet = workbook[holdings_sheet]
    found: dict[str, tuple[Decimal, str]] = {}
    for row in range(2, sheet.max_row + 1):
        label, cell = sheet.cell(row=row, column=1).value, sheet.cell(row=row, column=_VALUE_COLUMN)
        if not isinstance(label, str) or label == TOTAL_LABEL or cell.comment is None:
            continue
        # Excel may put the author's name before the text, so the text is looked for, not matched.
        kind = next((kind for marker, kind in _COMMENT_KINDS if marker in (cell.comment.text or "")), None)
        value = cell.value
        if kind and isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0:
            found[_ticker(label)] = (Decimal(str(value)), kind)
    return found


class ExcelExporter:
    """Writes "Current Holdings" and "Portfolio History" into an .xlsx workbook."""

    def __init__(
        self,
        path: Path,
        holdings_sheet: str = DEFAULT_HOLDINGS_SHEET,
        history_sheet: str = DEFAULT_HISTORY_SHEET,
        index_names: Mapping[str, str] | None = None,
        backups_kept: int = BACKUPS_KEPT,
    ) -> None:
        self.path = path
        self.backups_kept = max(0, backups_kept)
        """Backup copies to keep in all, at most one a day. 0 makes no copy, and leaves copies made earlier alone."""
        self.index_names = dict(index_names or {})
        """Names to show next to tickers in the Investment column, e.g. "(ABCD) World fund"."""
        self.holdings_sheet = holdings_sheet
        self.history_sheet = history_sheet

    def export(self, valuation: PortfolioValuation, today: date) -> ExportResult:
        """Update the workbook, creating it (or our sheets) if missing."""
        if self.path.suffix.lower() != ".xlsx":
            raise UserFacingError(
                f'"{self.path.name}" is not an .xlsx spreadsheet.',
                "Choose an Excel workbook whose name ends in .xlsx (not .xls, .xlsm or .csv).",
            )
        workbook, created = self._open()

        holdings = self._sheet(workbook, self.holdings_sheet, HOLDINGS_HEADER, _EARLIER_HOLDINGS_HEADERS)
        history = self._sheet(workbook, self.history_sheet, HISTORY_HEADER)
        if created:
            for name in [name for name in workbook.sheetnames if name not in (self.holdings_sheet, self.history_sheet)]:
                del workbook[name]

        self._freeze_history(history, holdings)
        written, shown = self._write_holdings(holdings, valuation)
        # A history row is dated today, so it is only written when every value is today's: a price
        # typed in today counts, a last known one or one typed in on an earlier day does not.
        current = valuation.prices_are_from(today)
        history_action, history_shown = (
            self._write_history(history, valuation, today, written) if current else ("skipped", {})
        )
        _style(holdings, _HOLDINGS_FORMATS, bold_last_row=written > 0, shown=shown)
        _mark_missing_values(holdings, written + 1)
        _style(history, _HISTORY_FORMATS, shown=self._history_formulas(history, history_shown))
        self._mark_partial_history(history, written)

        notes: list[str] = []
        backup, made_now = (None, False) if created or not self.backups_kept else self._backup(today)
        if backup is not None:
            notes.append(f'A copy of the spreadsheet as it was before today\'s first update is in "{backup.name}".')
        if valuation.total_market_value is None:
            notes.append(
                f'Not every current value is available, so no row was written to "{self.history_sheet}" '
                "this time. Existing history was not changed."
            )
        elif valuation.uses_last_known_prices:
            notes.append(
                "Some values use a price from an earlier day because Yahoo could not be reached, so no row "
                f'was written to "{self.history_sheet}" this time. Existing history was not changed.'
            )
        elif not current:
            notes.append(
                "Some values were entered by you on an earlier day, so no row was written to "
                f'"{self.history_sheet}" this time. Existing history was not changed.'
            )
        elif valuation.left_out:
            tickers = ", ".join(holding.position.ticker for holding in valuation.left_out)
            notes.append(
                f'Today\'s row in "{self.history_sheet}" leaves out what has no price ({tickers}): its totals, '
                "Net Worth and Gain vs Deposits cover the rest only, until a value is typed in for it in "
                f'"{self.holdings_sheet}".'
            )
        try:
            self._save(workbook)
        except UserFacingError:
            if backup is not None and made_now:
                backup.unlink(missing_ok=True)  # the file was not changed, so the copy is not needed
            raise
        if made_now:
            self._drop_old_backups()
        return ExportResult(str(self.path), created, written, history_action, tuple(notes))

    # -- workbook handling -------------------------------------------------------

    def _open(self) -> tuple[Any, bool]:
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        if not self.path.exists():
            if not self.path.parent.exists():
                raise UserFacingError(
                    f'The folder "{self.path.parent}" does not exist.',
                    "Choose a folder that exists for the spreadsheet.",
                )
            return Workbook(), True
        try:
            return load_workbook(self.path), False
        except PermissionError as exc:
            raise self._locked(exc) from exc
        except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as exc:
            raise UserFacingError(
                f'"{self.path.name}" could not be opened as an Excel workbook.',
                "Check that it is a real .xlsx file and not damaged, or choose another file.",
                details=repr(exc),
            ) from exc

    def _sheet(
        self,
        workbook: Any,
        name: str,
        header: tuple[str, ...],
        earlier_headers: tuple[tuple[str, ...], ...] = (),
    ) -> Any:
        """Return our sheet, creating it if missing. Refuse a sheet that is not ours."""
        if name not in workbook.sheetnames:
            sheet = workbook.create_sheet(name)
            sheet.append(header)
            return sheet
        sheet = workbook[name]
        found = tuple(cell.value for cell in sheet[1])
        while found and found[-1] is None:
            found = found[:-1]
        if sheet.max_row == 1 and not found:
            for column, title in enumerate(header, start=1):
                sheet.cell(row=1, column=column, value=title)
        elif found != header and found not in earlier_headers:
            raise UserFacingError(
                f'The spreadsheet already has a sheet called "{name}", but it was not made by this app.',
                "Rename that sheet in Excel, or choose another spreadsheet. Nothing was changed.",
            )
        return sheet

    def _backup(self, today: date) -> tuple[Path, bool]:
        """Copy the workbook to a file next to it named after today, unless today's copy exists.

        That copy is the file as it was before the day's first update; later updates that day
        keep it. Returns the copy and whether it was made now.
        """
        backup = self.path.with_name(f"{self.path.stem}_backup_{today.isoformat()}{self.path.suffix}")
        if backup.exists():
            return backup, False
        try:
            shutil.copy2(self.path, backup)
        except PermissionError as exc:
            raise self._locked(exc) from exc
        except OSError as exc:
            raise UserFacingError(
                f'A backup copy of "{self.path.name}" could not be made next to it.',
                "Check that the folder is not read-only and has free space, then try again. Nothing was changed.",
                details=repr(exc),
            ) from exc
        return backup, True

    def _drop_old_backups(self) -> None:
        for old in self.backups()[: -self.backups_kept]:
            try:
                old.unlink()
            except OSError:
                pass  # an old copy that cannot be deleted does no harm

    def backups(self) -> list[Path]:
        """This workbook's dated backup copies, oldest first."""
        pattern = re.compile(rf"{re.escape(self.path.stem)}_backup_\d{{4}}-\d\d-\d\d{re.escape(self.path.suffix)}")
        found = [file for file in self.path.parent.iterdir() if pattern.fullmatch(file.name) and file.is_file()]
        return sorted(found, key=lambda file: file.name)

    def _locked(self, exc: Exception) -> UserFacingError:
        return UserFacingError(
            f'"{self.path.name}" could not be saved because it is open in another program.',
            "Close the spreadsheet in Excel and try again. Nothing was changed.",
            details=repr(exc),
        )

    def _save(self, workbook: Any) -> None:
        # Save beside the target and swap it in, so a failed save never leaves half a file.
        temporary = self.path.with_name(f"~{self.path.stem}.saving{self.path.suffix}")
        try:
            workbook.save(temporary)
            os.replace(temporary, self.path)
        except PermissionError as exc:
            raise self._locked(exc) from exc
        except OSError as exc:
            raise UserFacingError(
                f'"{self.path.name}" could not be saved.',
                "Check that the folder still exists and that there is free disk space, then try again.",
                details=repr(exc),
            ) from exc
        finally:
            temporary.unlink(missing_ok=True)

    # -- sheet contents ----------------------------------------------------------

    def _write_holdings(self, sheet: Any, valuation: PortfolioValuation) -> tuple[int, dict[tuple[int, int], Any]]:
        """Rewrite the whole table. The sheet itself stays, so formulas pointing at it keep working.

        The Current Value of each investment is a number, or "unavailable" for the user to type a
        value over. Everything worked out from those values is an Excel formula, so the shares, the
        gains and the Total row follow what is typed, as the app's own table does: an investment
        without a value is left out of the totals. Returns the rows written, and for each formula
        cell the value it shows now, to size the columns by.
        """
        if sheet.max_column > len(HOLDINGS_HEADER):
            sheet.delete_cols(len(HOLDINGS_HEADER) + 1, sheet.max_column - len(HOLDINGS_HEADER))
        sheet.delete_rows(1, sheet.max_row)
        rows: list[list[Any]] = [list(HOLDINGS_HEADER)]
        shown: dict[tuple[int, int], Any] = {}
        currency = valuation.portfolio.currency
        total = len(valuation.holdings) + 2
        values = f"D2:D{total - 1}"
        # With no value at all, nothing is left out: the totals are simply unavailable.
        none_yet = f"COUNT({values})=0"

        def in_totals(column: str) -> str:
            """The sum of ``column`` over the investments the totals cover."""
            span = f"{column}2:{column}{total - 1}"
            return f"IF({none_yet},SUM({span}),SUMPRODUCT({span},--ISNUMBER({values})))"

        for row, holding in enumerate(valuation.holdings, start=2):
            position = holding.position
            name = self.index_names.get(position.ticker)
            rows.append(
                [
                    f"({position.ticker}) {name}" if name else position.ticker,
                    position.cost_basis,
                    f'=IF(OR(ISNUMBER(D{row}),{none_yet}),B{row}/B${total},"{NOT_IN_TOTAL}")',
                    _cell(holding.market_value),
                    f'=IF(ISNUMBER(D{row}),D{row}/D${total},IF({none_yet},"{UNAVAILABLE}","{NOT_IN_TOTAL}"))',
                    f'=IF(ISNUMBER(D{row}),D{row}/B{row}-1,"{UNAVAILABLE}")',
                    position.dividends,
                    currency,
                ]
            )
            shown[row, 3] = _percent(holding.invested_percent) if holding.in_totals else NOT_IN_TOTAL
            shown[row, 5] = _percent(holding.allocation_percent) if holding.in_totals else NOT_IN_TOTAL
            shown[row, 6] = _percent(holding.percent_change)
        if valuation.holdings:
            invested = valuation.invested_in_totals
            unrealized = valuation.unrealized_pnl
            rows.append(
                [
                    TOTAL_LABEL,
                    f"={in_totals('B')}",
                    f'=IF(B{total}>0,1,"{UNAVAILABLE}")',
                    f'=IF({none_yet},"{UNAVAILABLE}",SUM({values}))',
                    f'=IF(ISNUMBER(D{total}),1,"{UNAVAILABLE}")',
                    f'=IF(ISNUMBER(D{total}),D{total}/B{total}-1,"{UNAVAILABLE}")',
                    f"={in_totals('G')}",
                    currency,
                ]
            )
            shown[total, 2] = invested
            shown[total, 3] = _percent(_HUNDRED if invested > 0 else None)
            shown[total, 4] = _cell(valuation.total_market_value)
            shown[total, 5] = _percent(_HUNDRED if valuation.total_market_value else None)
            shown[total, 6] = _cell(unrealized / invested if unrealized is not None and invested > 0 else None)
            shown[total, 7] = sum(
                (holding.position.dividends for holding in valuation.holdings if holding.in_totals), Decimal(0)
            )
        for row_number, row_values in enumerate(rows, start=1):
            for column, value in enumerate(row_values, start=1):
                sheet.cell(row=row_number, column=column, value=value)

        # A value at a price from an earlier run, or at one the user typed in, gets a comment (Excel
        # shows a red corner) saying which and from when. A missing one says it can be typed in.
        for row, holding in enumerate(valuation.holdings, start=2):
            price = holding.price
            if price is not None and price.last_known:
                day = price.as_of.date().isoformat() if price.as_of else "an earlier day"
                sheet.cell(row=row, column=_VALUE_COLUMN).comment = _comment(
                    f"{LAST_KNOWN_COMMENT}, from {day}. Yahoo could not be reached when this was written."
                )
            elif price is not None and price.entered:
                day = price.as_of.astimezone().date().isoformat() if price.as_of else "an earlier day"
                sheet.cell(row=row, column=_VALUE_COLUMN).comment = _comment(
                    f"{ENTERED_COMMENT} on {day}, not from Yahoo."
                )
            elif price is None:
                sheet.cell(row=row, column=_VALUE_COLUMN).comment = _comment(TYPE_A_VALUE)
        if valuation.total_market_value is not None:
            total_notes = []
            if valuation.uses_last_known_prices:
                total_notes.append("Includes last known prices from an earlier day, so this total is not up to date.")
            if valuation.uses_entered_prices:
                total_notes.append("Includes values entered by hand.")
            if total_notes:
                sheet.cell(row=total, column=_VALUE_COLUMN).comment = _comment(" ".join(total_notes))
        return len(valuation.holdings), shown

    # The history columns, by letter: A Date, B Total Cost Basis, C Total Market Value, D Unrealized
    # PnL, E Realized PnL, F ROI %, G Net Worth, H Net Deposits, I Gain vs Deposits. Only B, C, E and
    # H are figures; D, F, G and I are formulas over them, so they follow when a figure changes. Net
    # Worth is the market value plus the cash the statement leaves, e.g. "C5+504.76".

    def _write_history(
        self, sheet: Any, valuation: PortfolioValuation, today: date, holdings_rows: int
    ) -> tuple[str, dict[tuple[int, int], Any]]:
        """Add today's row, or update it if today already has one. Other rows keep their figures.

        Today's cost basis and market value point at the Total row of the holdings sheet, so a
        value typed in there reaches the history too; ``_freeze_history`` turns them back into
        figures on a later run. Returns what the row's formula cells show, to size columns by.
        """
        portfolio = valuation.portfolio
        if valuation.total_market_value is None or valuation.unrealized_pnl is None:
            return "skipped", {}
        target = None
        for row in range(2, sheet.max_row + 1):
            existing = sheet.cell(row=row, column=1).value
            if isinstance(existing, datetime):
                existing = existing.date()
            if existing == today or existing == today.isoformat():
                target = row
                break
        action = "updated" if target else "added"
        row = target or sheet.max_row + 1

        if holdings_rows:
            holdings = _reference(self.holdings_sheet)
            total = holdings_rows + 2
            cost_basis, market_value = f"={holdings}B{total}", f"={holdings}D{total}"
        else:
            cost_basis, market_value = valuation.invested_in_totals, valuation.total_market_value
        # Like the Total row, net worth covers the investments that have a value.
        net_worth = f'=IF(ISNUMBER(C{row}),C{row}{_signed(portfolio.cash)},"{UNAVAILABLE}")'
        values = [today, cost_basis, market_value, None, portfolio.realized_pnl, None, net_worth, portfolio.net_deposits]
        for column, value in enumerate(values, start=1):
            if value is not None:
                sheet.cell(row=row, column=column, value=value)
        left_out = [holding.position.ticker for holding in valuation.left_out]
        sheet.cell(row=row, column=1).comment = (
            _comment(_partial_comment(left_out, self.holdings_sheet, live=bool(holdings_rows)), height=130)
            if left_out
            else None
        )
        return action, {
            (row, 2): valuation.invested_in_totals,
            (row, 3): valuation.total_market_value,
            (row, 7): _cell(valuation.net_worth),
        }

    def _freeze_history(self, history: Any, holdings: Any) -> None:
        """Give history rows that point at the holdings sheet the figures it shows now, before it is
        rewritten: each row keeps what it showed on its day, including values typed in since."""
        linked = [row for row in range(2, history.max_row + 1) if _points_elsewhere(history.cell(row=row, column=2).value)]
        if not linked:
            return
        cost_basis, market_value, missing = _holdings_totals(holdings)
        for row in linked:
            history.cell(row=row, column=2).value = cost_basis
            history.cell(row=row, column=3).value = _cell(market_value)
            history.cell(row=row, column=1).comment = (
                _comment(_partial_comment(missing, self.holdings_sheet, live=False), height=110) if missing else None
            )

    def _mark_partial_history(self, sheet: Any, holdings_rows: int) -> None:
        """Colour the history rows that leave out investments without a value, the rows with our
        comment. A row still pointing at the holdings sheet is coloured only while a value is
        missing there, so it turns plain once the value is typed in. Replaces the sheet's rules."""
        from openpyxl.formatting.formatting import ConditionalFormattingList
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Font, PatternFill

        sheet.conditional_formatting = ConditionalFormattingList()  # the app's sheet: no rules but its own
        fill = PatternFill(fill_type="solid", start_color=_PARTIAL_FILL_COLOUR, end_color=_PARTIAL_FILL_COLOUR)
        font = Font(color=_PARTIAL_TEXT_COLOUR)
        # INDIRECT, because Google Sheets does not let a rule name another sheet directly.
        values = f"{_reference(self.holdings_sheet)}D2:D{holdings_rows + 1}".replace('"', '""')
        for row in range(2, sheet.max_row + 1):
            comment = sheet.cell(row=row, column=1).comment
            # Excel may put the author's name before the text, so the text is looked for, not matched.
            if comment is None or PARTIAL_COMMENT not in (comment.text or ""):
                continue
            linked = _points_elsewhere(sheet.cell(row=row, column=2).value) and holdings_rows
            condition = f'COUNT(INDIRECT("{values}"))<{holdings_rows}' if linked else "TRUE"
            sheet.conditional_formatting.add(f"A{row}:I{row}", FormulaRule(formula=[condition], fill=fill, font=font))

    def _history_formulas(self, sheet: Any, shown: dict[tuple[int, int], Any]) -> dict[tuple[int, int], Any]:
        """Write the formulas of every dated history row and return what they show, with ``shown``.

        Rows written by earlier versions hold figures in every column, and an ROI % that was the
        unrealized PnL over the cost basis; they get the same formulas, which give the same figures
        except for ROI %, now the gain over the money deposited.
        """
        shown = dict(shown)
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=1).value is None:
                continue

            def number(column: int) -> Decimal | None:
                value = shown.get((row, column), sheet.cell(row=row, column=column).value)
                return Decimal(str(value)) if _is_number(value) else None

            cost_basis, market_value, deposits = number(2), number(3), number(8)
            net_worth = sheet.cell(row=row, column=7).value
            if _is_number(net_worth) and market_value is not None:  # a figure, from an earlier version
                cash = _signed(Decimal(str(net_worth)) - market_value)
                sheet.cell(row=row, column=7).value = f'=IF(ISNUMBER(C{row}),C{row}{cash},"{UNAVAILABLE}")'
                shown[row, 7] = Decimal(str(net_worth))
            elif (row, 7) not in shown:
                cash = _NET_WORTH_CASH.search(str(net_worth))
                worth = market_value + Decimal(cash.group(1)) if cash and market_value is not None else None
                shown[row, 7] = _cell(worth)
            net_worth_now = number(7)

            sheet.cell(row=row, column=4).value = f'=IF(AND(ISNUMBER(C{row}),ISNUMBER(B{row})),C{row}-B{row},"{UNAVAILABLE}")'
            sheet.cell(row=row, column=6).value = f'=IF(AND(ISNUMBER(I{row}),N(H{row})>0),I{row}/H{row},"{UNAVAILABLE}")'
            sheet.cell(row=row, column=9).value = f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(H{row})),G{row}-H{row},"{UNAVAILABLE}")'
            gain = net_worth_now - deposits if net_worth_now is not None and deposits is not None else None
            shown[row, 4] = _cell(market_value - cost_basis if market_value is not None and cost_basis is not None else None)
            shown[row, 6] = _cell(gain / deposits if gain is not None and deposits else None)
            shown[row, 9] = _cell(gain)
        return shown
